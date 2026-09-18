import base64
import json
import struct
import zipfile
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import fitz
import pytest
from docx import Document
from openpyxl import Workbook

from pkas.config import Settings
from pkas.document_extraction import DocumentExtractionPipeline
from pkas.extraction_eval import text_metrics
from pkas.ingest import chunk_document
from pkas.parsers import ParsedBlock, ParsedDocument, ParseError, parse_file
from pkas.system import KnowledgeSystem


def test_jsonl_message_blocks_preserve_time_and_speaker(
    test_settings: Settings,
    source_root: Path,
) -> None:
    path = source_root / "messages.jsonl"
    records = [
        {
            "sent_at": "2026-08-07 10:00:00",
            "speaker": "客户甲",
            "text": "需要更新项目排期",
        }
    ]
    path.write_text(
        "\n".join(json.dumps(record, ensure_ascii=False) for record in records),
        encoding="utf-8",
    )

    parsed = parse_file(path, settings=test_settings)
    chunks = chunk_document(parsed, max_chars=200)

    expected = "[2026-08-07 10:00:00 · 客户甲] 需要更新项目排期"
    assert parsed.blocks[0].text == expected
    assert chunks[0]["text"] == expected


def test_invalid_json_falls_back_to_readable_text_with_quality_warning(
    test_settings: Settings,
    source_root: Path,
) -> None:
    path = source_root / "incomplete.json"
    path.write_text('{"project":"knowledge-base",\n"state":"unfinished",', encoding="utf-8")

    parsed = parse_file(path, settings=test_settings)

    assert parsed.parser_name == "invalid-json-text-fallback"
    assert "knowledge-base" in parsed.text
    assert parsed.metadata["json_structure_valid"] is False
    assert parsed.metadata["json_structure_not_read"] is True
    reasons = parsed.metadata["extraction"]["initial_quality"]["reasons"]
    assert "invalid_json_text_fallback" in reasons


def test_pdf_records_low_text_pages_and_keeps_native_provenance(
    test_settings: Settings,
    source_root: Path,
) -> None:
    path = source_root / "mixed.pdf"
    document = fitz.open()
    first = document.new_page()
    first.insert_text((72, 72), "native searchable project evidence")
    document.new_page()
    document.save(path)
    document.close()

    parsed = parse_file(path, settings=test_settings)

    extraction = parsed.metadata["extraction"]
    assert parsed.parser_name == "pymupdf-blocks"
    assert "native searchable project evidence" in parsed.text
    assert extraction["initial_quality"]["visual_pages"] == [2]
    assert extraction["warnings"] == ["paddleocr:remote_processing_not_enabled"]
    assert parsed.blocks[0].locator.startswith("page:1/block:")


def test_pdf_uses_local_poppler_fallback_only_when_native_has_no_body_text(
    test_settings: Settings,
    source_root: Path,
    monkeypatch,
) -> None:
    path = source_root / "fallback.pdf"
    document = fitz.open()
    document.new_page()
    document.save(path)
    document.close()
    calls: list[list[str]] = []

    monkeypatch.setattr(
        "pkas.parsers.shutil.which",
        lambda name: "pdftotext" if name == "pdftotext" else None,
    )

    def fake_run(command, **_kwargs):
        calls.append(command)
        return SimpleNamespace(
            returncode=0,
            stdout="备用本地 PDF 文字\f第二页文字".encode(),
        )

    monkeypatch.setattr("pkas.parsers.subprocess.run", fake_run)
    parsed = parse_file(path, settings=test_settings)

    assert "备用本地 PDF 文字" in parsed.text
    assert parsed.parser_name == "pymupdf-blocks+poppler-pdftotext"
    assert parsed.metadata["poppler_fallback"] == "used"
    assert parsed.blocks[0].locator == "page:1/poppler-line:1"
    assert calls == [["pdftotext", "-layout", str(path), "-"]]


def test_pdf_uses_poppler_only_for_pages_missing_native_text(
    test_settings: Settings,
    source_root: Path,
    monkeypatch,
) -> None:
    path = source_root / "mixed-fallback.pdf"
    document = fitz.open()
    first = document.new_page()
    first.insert_text((72, 72), "native first page evidence")
    document.new_page()
    document.save(path)
    document.close()
    calls: list[list[str]] = []
    monkeypatch.setattr(
        "pkas.parsers.shutil.which",
        lambda name: "pdftotext" if name == "pdftotext" else None,
    )

    def fake_run(command, **_kwargs):
        calls.append(command)
        return SimpleNamespace(
            returncode=0,
            stdout=b"native first page evidence\fRecovered second page evidence",
        )

    monkeypatch.setattr("pkas.parsers.subprocess.run", fake_run)
    parsed = parse_file(path, settings=test_settings)

    assert parsed.parser_name == "pymupdf-blocks+poppler-pdftotext"
    assert parsed.metadata["missing_native_text_pages"] == [2]
    assert parsed.metadata["poppler_fallback"] == "used_on_missing_native_pages"
    assert "Recovered second page evidence" in parsed.text
    assert "native first page evidence" in parsed.text
    assert all(block.page == 2 for block in parsed.blocks if "Recovered" in block.text)
    assert calls == [["pdftotext", "-layout", str(path), "-"]]


def test_standalone_image_keeps_local_metadata_without_claiming_pixel_understanding(
    test_settings: Settings,
    source_root: Path,
) -> None:
    path = source_root / "architecture.png"
    path.write_bytes(
        b"\x89PNG\r\n\x1a\n" + b"\x00\x00\x00\rIHDR" + struct.pack(">II", 1280, 720)
    )

    parsed = parse_file(path, settings=test_settings)

    assert parsed.parser_name == "local-image-metadata"
    assert parsed.metadata["image_width"] == 1280
    assert parsed.metadata["image_height"] == 720
    assert parsed.metadata["pixel_content_not_read"] is True
    assert "尺寸 1280 × 720" in parsed.text
    extraction = parsed.metadata["extraction"]
    assert extraction["initial_quality"]["score"] == 0.15
    assert "visual_pixels_not_read" in extraction["initial_quality"]["reasons"]
    assert extraction["warnings"] == ["paddleocr:remote_processing_not_enabled"]


def test_standalone_png_indexes_authored_local_description_but_not_pixels(
    test_settings: Settings,
    source_root: Path,
) -> None:
    path = source_root / "annotated-architecture.png"
    title = b"Knowledge Base Delivery Architecture"
    description = "本地图片说明，可用于检索。".encode()
    text_chunk = b"Title\x00" + title
    itxt_chunk = b"Description\x00\x00\x00\x00" + description
    png = b"\x89PNG\r\n\x1a\n" + b"\x00\x00\x00\rIHDR" + struct.pack(
        ">II", 1280, 720
    ) + b"\x08\x02\x00\x00\x00\x00\x00\x00\x00"
    png += struct.pack(
        ">I", len(text_chunk)
    ) + b"tEXt" + text_chunk + b"\x00\x00\x00\x00"
    png += struct.pack(
        ">I", len(itxt_chunk)
    ) + b"iTXt" + itxt_chunk + b"\x00\x00\x00\x00"
    png += b"\x00\x00\x00\x00IEND\xaeB`\x82"
    path.write_bytes(png)

    parsed = parse_file(path, settings=test_settings)

    assert parsed.metadata["local_description_count"] == 2
    assert "Knowledge Base Delivery Architecture" in parsed.text
    assert "本地图片说明，可用于检索。" in parsed.text
    assert any(block.kind == "image-description" for block in parsed.blocks)
    assert parsed.metadata["pixel_content_not_read"] is True


def test_pdf_suppresses_repeated_short_headers_but_keeps_page_body(
    test_settings: Settings,
    source_root: Path,
) -> None:
    path = source_root / "repeated-header.pdf"
    document = fitz.open()
    for page_number in range(1, 5):
        page = document.new_page()
        page.insert_text((72, 40), "CONFIDENTIAL PROJECT HEADER")
        page.insert_text((72, 160), f"Unique body evidence page {page_number}")
    document.save(path)
    document.close()

    parsed = parse_file(path, settings=test_settings)

    assert "CONFIDENTIAL PROJECT HEADER" not in parsed.text
    assert all(f"Unique body evidence page {page}" in parsed.text for page in range(1, 5))
    assert parsed.metadata["suppressed_repeated_margin_blocks"] == 4


def test_pdf_parser_keeps_outline_annotations_and_hyperlinks(
    test_settings: Settings,
    source_root: Path,
) -> None:
    path = source_root / "review.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "项目评审正文")
    page.add_text_annot((120, 120), "需要补充验收说明")
    page.insert_link(
        {
            "kind": fitz.LINK_URI,
            "from": fitz.Rect(72, 72, 210, 90),
            "uri": "https://example.com/review",
        }
    )
    document.set_toc([[1, "项目评审目录", 1]])
    document.save(path)
    document.close()

    parsed = parse_file(path, settings=test_settings)

    assert any(block.kind == "heading" and block.text == "项目评审目录" for block in parsed.blocks)
    assert any("需要补充验收说明" in block.text for block in parsed.blocks)
    assert any("https://example.com/review" in block.text for block in parsed.blocks)
    assert parsed.metadata["outline_count"] == 1
    assert parsed.metadata["annotation_count"] == 1
    assert parsed.metadata["hyperlink_count"] == 1


def test_pdf_parser_indexes_safe_filled_form_fields_and_redacts_password_fields(
    test_settings: Settings,
    source_root: Path,
) -> None:
    path = source_root / "filled-form.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "报价登记表")
    customer = fitz.Widget()
    customer.field_name = "客户名称"
    customer.field_value = "示例客户"
    customer.field_type = fitz.PDF_WIDGET_TYPE_TEXT
    customer.rect = fitz.Rect(72, 100, 260, 124)
    page.add_widget(customer)
    password = fitz.Widget()
    password.field_name = "password"
    password.field_value = "never-index-this"
    password.field_type = fitz.PDF_WIDGET_TYPE_TEXT
    password.rect = fitz.Rect(72, 136, 260, 160)
    page.add_widget(password)
    document.save(path)
    document.close()

    parsed = parse_file(path, settings=test_settings)

    assert "PDF 表单：字段 客户名称；类型 Text；值 示例客户" in parsed.text
    assert "never-index-this" not in parsed.text
    assert parsed.metadata["form_field_count"] == 1
    assert parsed.metadata["redacted_form_field_count"] == 1
    assert any(block.kind == "pdf-form-field" for block in parsed.blocks)


def test_paddleocr_service_fills_scanned_pdf_page_when_explicitly_enabled(
    test_settings: Settings,
    source_root: Path,
) -> None:
    path = source_root / "scan.pdf"
    document = fitz.open()
    document.new_page()
    document.save(path)
    document.close()
    settings = Settings(
        project_root=test_settings.project_root,
        data_root=test_settings.data_root,
        allowed_origins=[],
        document_paddleocr_enabled=True,
        document_ai_enhancement_enabled=True,
        document_paddleocr_base_url="http://127.0.0.1:8080",
        document_allow_remote_processing=True,
    )
    payloads: list[dict[str, Any]] = []

    def transport(payload: dict[str, Any]) -> dict[str, Any]:
        payloads.append(payload)
        return {
            "result": {
                "layoutParsingResults": [
                    {
                        "prunedResult": {
                            "parsing_res_list": [
                                {
                                    "block_label": "text",
                                    "block_content": "扫描页识别文字",
                                    "block_bbox": [10, 20, 200, 80],
                                }
                            ]
                        }
                    }
                ]
            }
        }

    parsed = parse_file(path, settings=settings, paddle_transport=transport)

    assert parsed.text == "扫描页识别文字"
    assert "paddleocr-vl-1.6-service" in parsed.parser_name
    assert parsed.blocks[0].locator == "page:1/visual-block:1"
    assert parsed.blocks[0].bbox == (10.0, 20.0, 200.0, 80.0)
    assert payloads[0]["fileType"] == 1
    assert payloads[0]["visualize"] is False


def test_restricted_document_is_not_sent_to_remote_parser(
    test_settings: Settings,
    source_root: Path,
) -> None:
    path = source_root / "restricted-scan.pdf"
    document = fitz.open()
    document.new_page()
    document.save(path)
    document.close()
    settings = Settings(
        project_root=test_settings.project_root,
        data_root=test_settings.data_root,
        allowed_origins=[],
        document_paddleocr_enabled=True,
        document_ai_enhancement_enabled=True,
        document_paddleocr_base_url="http://127.0.0.1:8080",
        document_allow_remote_processing=True,
    )
    called = False

    def transport(_payload: dict[str, Any]) -> dict[str, Any]:
        nonlocal called
        called = True
        return {}

    with pytest.raises(ParseError, match="没有可索引文字"):
        parse_file(
            path,
            settings=settings,
            privacy="restricted",
            paddle_transport=transport,
        )
    assert called is False


def test_pptx_openxml_parser_extracts_slide_table_and_notes(
    test_settings: Settings,
    source_root: Path,
) -> None:
    path = source_root / "project.pptx"
    slide_xml = """<?xml version="1.0" encoding="UTF-8"?>
    <p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
           xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
      <p:cSld><p:spTree>
        <p:sp><p:txBody><a:p><a:r><a:t>项目总览</a:t></a:r></a:p></p:txBody></p:sp>
        <a:tbl><a:tr><a:tc><a:txBody><a:p><a:r><a:t>关键指标</a:t></a:r></a:p></a:txBody></a:tc></a:tr></a:tbl>
      </p:spTree></p:cSld>
    </p:sld>"""
    notes_xml = """<?xml version="1.0" encoding="UTF-8"?>
    <p:notes xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
             xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
      <p:cSld><p:spTree><p:sp><p:txBody><a:p><a:r><a:t>讲解备注</a:t></a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld>
    </p:notes>"""
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("ppt/slides/slide1.xml", slide_xml)
        archive.writestr("ppt/notesSlides/notesSlide1.xml", notes_xml)

    parsed = parse_file(path, settings=test_settings)

    assert "项目总览" in parsed.text
    assert "关键指标" in parsed.text
    assert "讲解备注" in parsed.text
    assert parsed.metadata["slide_count"] == 1
    assert {block.kind for block in parsed.blocks} >= {"heading", "table-row", "speaker-note"}


def test_macro_and_template_openxml_variants_use_readonly_structured_parsers(
    test_settings: Settings,
    source_root: Path,
) -> None:
    document = Document()
    document.add_paragraph("宏文档仍按正文只读解析")
    docx = source_root / "brief.docx"
    document.save(docx)
    docm = source_root / "brief.docm"
    docm.write_bytes(docx.read_bytes())

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["模板表头", "值"])
    sheet.append(["阶段", "验收"])
    xlsx = source_root / "ledger.xlsx"
    workbook.save(xlsx)
    xltm = source_root / "ledger.xltm"
    xltm.write_bytes(xlsx.read_bytes())

    pptx = source_root / "review.pptx"
    with zipfile.ZipFile(pptx, "w") as archive:
        archive.writestr(
            "ppt/slides/slide1.xml",
            """<p:sld xmlns:p=\"http://schemas.openxmlformats.org/presentationml/2006/main\"
            xmlns:a=\"http://schemas.openxmlformats.org/drawingml/2006/main\"><p:cSld><p:spTree>
            <p:sp><p:txBody><a:p><a:r><a:t>模板演示正文</a:t></a:r></a:p></p:txBody></p:sp>
            </p:spTree></p:cSld></p:sld>""",
        )
    ppsm = source_root / "review.ppsm"
    ppsm.write_bytes(pptx.read_bytes())

    parsed_doc = parse_file(docm, settings=test_settings)
    parsed_sheet = parse_file(xltm, settings=test_settings)
    parsed_slides = parse_file(ppsm, settings=test_settings)

    assert parsed_doc.parser_name == "python-docx-structured"
    assert "宏文档仍按正文只读解析" in parsed_doc.text
    assert parsed_sheet.parser_name == "openpyxl-structured"
    assert "A1=模板表头" in parsed_sheet.text
    assert parsed_slides.parser_name == "openxml-pptx-structured"
    assert "模板演示正文" in parsed_slides.text


def test_xlsx_parser_preserves_formula_and_cell_coordinates(
    test_settings: Settings,
    source_root: Path,
) -> None:
    path = source_root / "inventory.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "项目台账"
    sheet["A1"] = "数量"
    sheet["A2"] = 2
    sheet["B2"] = "=A2*3"
    workbook.save(path)

    parsed = parse_file(path, settings=test_settings)

    assert "A1=数量" in parsed.text
    assert "B2=[公式 =A2*3]" in parsed.text
    assert parsed.metadata["formula_count"] == 1
    assert any(block.locator == "sheet:项目台账/row:2" for block in parsed.blocks)
    table_rows = [block for block in parsed.blocks if block.kind == "table-row"]
    assert table_rows[0].metadata["is_header"] is True
    assert table_rows[1].metadata["table_id"] == "sheet:项目台账"


def test_xlsx_parser_indexes_named_ranges_and_data_validation_rules(
    test_settings: Settings,
    source_root: Path,
) -> None:
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation

    path = source_root / "business-rules.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "项目台账"
    sheet.append(["客户", "状态", "金额"])
    sheet.append(["甲方", "进行中", 100])
    sheet.append(["乙方", "已交付", 200])
    workbook.defined_names.add(
        DefinedName("客户清单", attr_text="'项目台账'!$A$2:$A$3")
    )
    validation = DataValidation(
        type="list",
        formula1='"进行中,已交付"',
        allow_blank=False,
    )
    sheet.add_data_validation(validation)
    validation.add("B2:B100")
    workbook.save(path)

    parsed = parse_file(path, settings=test_settings)

    assert "Excel 命名区域：名称 客户清单" in parsed.text
    assert "引用 '项目台账'!$A$2:$A$3" in parsed.text
    assert "Excel 数据校验：工作表 项目台账；范围 B2:B100；类型 list" in parsed.text
    assert "formula1 \"进行中,已交付\"" in parsed.text
    assert parsed.metadata["named_range_count"] == 1
    assert parsed.metadata["data_validation_count"] == 1
    assert parsed.metadata["data_validation_truncated"] is False
    assert {block.kind for block in parsed.blocks} >= {"named-range", "data-validation"}


def test_xlsx_parser_indexes_chart_labels_and_data_references(
    test_settings: Settings,
    source_root: Path,
) -> None:
    from openpyxl.chart import BarChart, Reference

    path = source_root / "chart.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["月份", "销售额"])
    sheet.append(["一月", 120])
    sheet.append(["二月", 180])
    chart = BarChart()
    chart.title = "月度销售额"
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=3))
    sheet.add_chart(chart, "D2")
    workbook.save(path)

    parsed = parse_file(path, settings=test_settings)

    chart_block = next(block for block in parsed.blocks if block.kind == "chart-description")
    assert "月度销售额" in chart_block.text
    assert "销售额" in chart_block.text
    assert chart_block.metadata["pixel_content_not_read"] is True
    assert parsed.metadata["chart_count"] == 1
    assert parsed.metadata["chart_description_count"] == 1


def test_xlsx_parser_separates_logical_tables_and_carries_headers(
    test_settings: Settings,
    source_root: Path,
) -> None:
    path = source_root / "multi-table.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "总览"
    sheet.append(["项目交付台账"])
    sheet.append(["项目", "状态", "负责人"])
    sheet.append(["知识底座", "进行中", "王工"])
    sheet.append([])
    sheet.append(["客户跟进"])
    sheet.append(["客户", "当前事项"])
    sheet.append(["甲方", "等待确认"])
    workbook.save(path)

    parsed = parse_file(path, settings=test_settings)

    title = next(block for block in parsed.blocks if block.kind == "table-title")
    first_header = next(
        block
        for block in parsed.blocks
        if block.kind == "table-row" and block.metadata["is_header"]
    )
    project_row = next(block for block in parsed.blocks if "知识底座" in block.text)
    customer_header = next(
        block
        for block in parsed.blocks
        if block.kind == "table-row"
        and block.metadata["is_header"]
        and "客户" in block.text
    )

    assert title.text == "A1=项目交付台账"
    assert title.metadata["role"] == "title"
    assert first_header.metadata["column_headers"] == ["项目", "状态", "负责人"]
    assert project_row.metadata["header_locator"] == "sheet:总览/row:2"
    assert customer_header.metadata["table_id"] == "sheet:总览/table:2"
    assert parsed.metadata["table_count"] == 2


def test_xlsx_parser_carries_real_merged_cell_context_without_guessing(
    test_settings: Settings,
    source_root: Path,
) -> None:
    path = source_root / "merged-context.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "项目"
    sheet.append(["类别", "事项"])
    sheet.append(["软件工程", "接口联调"])
    sheet.append([None, "验收文档"])
    sheet.merge_cells("A2:A3")
    workbook.save(path)

    parsed = parse_file(path, settings=test_settings)
    continued = next(
        block for block in parsed.blocks if block.locator == "sheet:项目/row:3"
    )

    assert "A3=软件工程" in continued.text
    assert continued.metadata["merged_from"] == ["A3←A2"]
    assert parsed.metadata["merged_range_count"] == 1
    assert parsed.metadata["merged_cell_count"] == 1
    assert not parsed.metadata["merged_cell_mapping_truncated"]


def test_openxml_parser_indexes_authored_image_alt_text_locally(
    test_settings: Settings,
    source_root: Path,
) -> None:
    image = source_root / "diagram.png"
    image.write_bytes(
        base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/"
            "9O6e7wAAAABJRU5ErkJggg=="
        )
    )
    path = source_root / "delivery-with-diagram.docx"
    document = Document()
    document.add_paragraph("本周交付说明")
    shape = document.add_picture(str(image))
    shape._inline.docPr.set("descr", "知识库交付架构关系图")
    document.save(path)

    parsed = parse_file(path, settings=test_settings)

    image_block = next(block for block in parsed.blocks if block.kind == "image-description")
    extraction = parsed.metadata["extraction"]
    assert image_block.text == "图片说明：知识库交付架构关系图"
    assert image_block.metadata["local_only"] is True
    assert parsed.metadata["embedded_image_count"] == 1
    assert parsed.metadata["embedded_image_description_count"] == 1
    assert "embedded_images_not_ocrd" not in extraction["initial_quality"]["reasons"]


def test_openxml_parser_indexes_embedded_image_title_without_pixel_reading(
    test_settings: Settings,
    source_root: Path,
) -> None:
    image = source_root / "metadata-diagram.png"
    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/"
        "9O6e7wAAAABJRU5ErkJggg=="
    )
    title_chunk = b"Title\x00Embedded delivery architecture"
    iend = png.rfind(b"\x00\x00\x00\x00IEND")
    assert iend > 0
    image.write_bytes(
        png[:iend]
        + struct.pack(">I", len(title_chunk))
        + b"tEXt"
        + title_chunk
        + b"\x00\x00\x00\x00"
        + png[iend:]
    )
    path = source_root / "delivery-with-metadata-diagram.docx"
    document = Document()
    document.add_picture(str(image))
    document.save(path)

    parsed = parse_file(path, settings=test_settings)

    image_block = next(
        block
        for block in parsed.blocks
        if block.metadata.get("description_origin") == "embedded-image-metadata"
    )
    extraction = parsed.metadata["extraction"]
    assert image_block.text == "图片说明：Embedded delivery architecture"
    assert image_block.metadata["pixel_content_not_read"] is True
    assert parsed.metadata["embedded_image_count"] == 1
    assert parsed.metadata["embedded_image_description_count"] == 1
    assert "embedded_images_not_ocrd" not in extraction["initial_quality"]["reasons"]


@pytest.mark.parametrize(
    ("name", "content", "parser", "delimiter"),
    [
        ("export.csv", '名称,备注\n知识库,"包含换行\n的单元格"\n', "csv-structured", "comma"),
        ("export.tsv", "名称\t备注\n知识库\t结构化导出\n", "tsv-structured", "tab"),
    ],
)
def test_delimited_tables_preserve_header_metadata_and_quoted_cells(
    test_settings: Settings,
    source_root: Path,
    name: str,
    content: str,
    parser: str,
    delimiter: str,
) -> None:
    path = source_root / name
    path.write_text(content, encoding="utf-8", newline="")

    parsed = parse_file(path, settings=test_settings)

    assert parsed.parser_name == parser
    assert parsed.metadata["delimiter"] == delimiter
    assert parsed.metadata["row_count"] == 2
    assert parsed.blocks[0].metadata["is_header"] is True
    assert parsed.blocks[1].metadata["is_header"] is False
    assert parsed.blocks[1].metadata["table_id"] == "delimited:1"
    if name.endswith(".csv"):
        assert "包含换行 的单元格" in parsed.blocks[1].text


def test_rtf_parser_keeps_paragraphs_and_unicode_escapes(
    test_settings: Settings,
    source_root: Path,
) -> None:
    path = source_root / "brief.rtf"
    path.write_bytes(r"{\rtf1\ansi\uc1 \u39033?\u30446?\par \u20132?\u20184?}".encode("ascii"))

    parsed = parse_file(path, settings=test_settings)

    assert parsed.parser_name == "rtf-local-text"
    assert "项目" in parsed.text
    assert "交付" in parsed.text
    assert len(parsed.blocks) == 2


def test_opendocument_parsers_keep_text_tables_formulas_and_slide_locations(
    test_settings: Settings,
    source_root: Path,
) -> None:
    odt = source_root / "brief.odt"
    ods = source_root / "ledger.ods"
    odp = source_root / "review.odp"
    namespaces = " ".join(
        (
            'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"',
            'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"',
            'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0"',
            'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"',
        )
    )
    with zipfile.ZipFile(odt, "w") as archive:
        archive.writestr(
            "content.xml",
            f"""<office:document-content {namespaces}><office:body><office:text>
            <text:h>项目总览</text:h><text:p>交付边界</text:p>
            <table:table table:name="事项"><table:table-header-rows><table:table-row>
            <table:table-cell><text:p>名称</text:p></table:table-cell><table:table-cell><text:p>状态</text:p></table:table-cell>
            </table:table-row></table:table-header-rows><table:table-row>
            <table:table-cell><text:p>知识库</text:p></table:table-cell><table:table-cell><text:p>进行中</text:p></table:table-cell>
            </table:table-row></table:table></office:text></office:body></office:document-content>""",
        )
    with zipfile.ZipFile(ods, "w") as archive:
        ods_content = "".join(
            [
                f"<office:document-content {namespaces}><office:body><office:spreadsheet>",
                "<table:table table:name=\"项目台账\"><table:table-row>",
                "<table:table-cell><text:p>预算</text:p></table:table-cell>",
                "</table:table-row><table:table-row>",
                "<table:table-cell table:formula=\"of:=SUM([.A1:.A1])\">",
                "<text:p>100</text:p></table:table-cell></table:table-row>",
                "</table:table></office:spreadsheet></office:body></office:document-content>",
            ]
        )
        archive.writestr(
            "content.xml",
            ods_content,
        )
    with zipfile.ZipFile(odp, "w") as archive:
        archive.writestr(
            "content.xml",
            f"""<office:document-content {namespaces}><office:body><office:presentation>
            <draw:page><text:p>方案评审</text:p><text:p>下一步验收</text:p></draw:page>
            </office:presentation></office:body></office:document-content>""",
        )

    parsed_odt = parse_file(odt, settings=test_settings)
    parsed_ods = parse_file(ods, settings=test_settings)
    parsed_odp = parse_file(odp, settings=test_settings)

    assert parsed_odt.parser_name == "opendocument-text-structured"
    assert "项目总览" in parsed_odt.text and "知识库 | 进行中" in parsed_odt.text
    assert any(
        block.metadata.get("is_header")
        for block in parsed_odt.blocks
        if block.kind == "table-row"
    )
    assert parsed_ods.parser_name == "opendocument-sheet-structured"
    assert "[公式 of:=SUM([.A1:.A1])]" in parsed_ods.text
    assert parsed_odp.metadata["slide_count"] == 1
    assert parsed_odp.blocks[0].locator == "slide:1/block:1"


def test_legacy_office_prefers_local_converter_without_docling(
    test_settings: Settings,
    source_root: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = source_root / "legacy.doc"
    path.write_bytes(b"legacy fixture")
    pipeline = DocumentExtractionPipeline(settings=test_settings)

    class FakeOffice:
        def parse(self, source: Path, runtime_root: Path) -> ParsedDocument:
            assert source == path
            assert runtime_root == test_settings.data_root / "runtime" / "office-conversion"
            return ParsedDocument(
                title="legacy",
                text="本地转换后的正文",
                parser_name="libreoffice-headless+python-docx-structured",
                blocks=[
                    ParsedBlock(
                        kind="paragraph",
                        text="本地转换后的正文",
                        locator="docx:paragraph:1",
                    )
                ],
                metadata={"converted_from": ".doc", "converter_ephemeral": True},
            )

    def forbidden(*_args: object, **_kwargs: object) -> ParsedDocument:
        raise AssertionError("本地转换成功时不应再调用Docling")

    pipeline.office = FakeOffice()  # type: ignore[assignment]
    monkeypatch.setattr(pipeline.docling, "parse", forbidden)
    parsed = pipeline.parse(path)

    assert parsed.text == "本地转换后的正文"
    assert parsed.metadata["converted_from"] == ".doc"
    assert parsed.metadata["extraction"]["warnings"] == []


def test_semantic_chunking_preserves_block_locators() -> None:
    parsed = ParsedDocument(
        title="结构化文档",
        text="标题\n内容",
        parser_name="test",
        blocks=[
            ParsedBlock(kind="heading", text="标题", locator="page:1/block:1", page=1),
            ParsedBlock(kind="paragraph", text="内容", locator="page:1/block:2", page=1),
        ],
    )

    chunks = chunk_document(parsed, max_chars=100)

    assert chunks == [
        {
            "sequence": 0,
            "text": "标题\n\n内容",
            "locator": "page:1/block:1..page:1/block:2",
            "block_sequences": [0, 1],
            "chunk_kind": "document-section",
            "chunker_version": "typed-v2",
        }
    ]


def test_typed_chunking_uses_table_message_and_code_profiles() -> None:
    table = ParsedDocument(
        title="台账",
        text="",
        parser_name="test",
        blocks=[
            ParsedBlock(kind="heading", text="项目台账", locator="sheet:台账"),
            ParsedBlock(kind="table-row", text="编号 | 状态", locator="row:1"),
            *[
                ParsedBlock(kind="table-row", text=f"{index} | 完成", locator=f"row:{index}")
                for index in range(2, 47)
            ],
        ],
    )
    messages = ParsedDocument(
        title="会话",
        text="",
        parser_name="test",
        blocks=[
            ParsedBlock(kind="message", text=f"用户：消息 {index}", locator=f"record:{index}")
            for index in range(35)
        ],
    )
    code = ParsedDocument(
        title="service",
        text="",
        parser_name="plain-text",
        metadata={"source_extension": ".py"},
        blocks=[
            ParsedBlock(kind="paragraph", text=f"value_{index} = {index}", locator=f"line:{index}")
            for index in range(180)
        ],
    )

    table_chunks = chunk_document(table)
    message_chunks = chunk_document(messages)
    code_chunks = chunk_document(code)

    assert len(table_chunks) >= 3
    assert all(chunk["chunk_kind"] == "table-rows" for chunk in table_chunks)
    assert all("编号 | 状态" in chunk["text"] for chunk in table_chunks)
    assert len(message_chunks) == 2
    assert all(chunk["chunk_kind"] == "message-window" for chunk in message_chunks)
    assert all(len(chunk["block_sequences"]) <= 30 for chunk in message_chunks)
    assert len(code_chunks) >= 2
    assert all(chunk["chunk_kind"] == "code" for chunk in code_chunks)


def test_typed_v2_uses_each_table_own_header() -> None:
    parsed = ParsedDocument(
        title="多表格",
        text="",
        parser_name="test",
        blocks=[
            ParsedBlock(
                kind="table-row",
                text="项目 | 状态",
                locator="table:1/row:1",
                metadata={"table_id": "1", "is_header": True},
            ),
            ParsedBlock(
                kind="table-row",
                text="A | 完成",
                locator="table:1/row:2",
                metadata={"table_id": "1"},
            ),
            ParsedBlock(
                kind="table-row",
                text="设备 | 型号",
                locator="table:2/row:1",
                metadata={"table_id": "2", "is_header": True},
            ),
            ParsedBlock(
                kind="table-row",
                text="主控 | ESP32",
                locator="table:2/row:2",
                metadata={"table_id": "2"},
            ),
        ],
    )

    chunks = chunk_document(parsed, max_chars=30)

    assert any(
        "项目 | 状态" in item["text"] and "A | 完成" in item["text"] for item in chunks
    )
    assert any(
        "设备 | 型号" in item["text"] and "主控 | ESP32" in item["text"] for item in chunks
    )
    assert not any(
        "项目 | 状态" in item["text"] and "主控 | ESP32" in item["text"] for item in chunks
    )


def test_ingested_chunks_have_resolvable_block_relations(
    knowledge_system: KnowledgeSystem,
    source_root: Path,
) -> None:
    note = source_root / "block-model.md"
    note.write_text("# 架构\n\n第一段。\n第二段。", encoding="utf-8")

    imported = knowledge_system.ingestion.import_file(
        note,
        domain="work",
        privacy="private",
    )

    with knowledge_system.database.connect() as connection:
        blocks = connection.execute(
            "SELECT COUNT(*) FROM blocks WHERE document_id=?",
            (imported["document_id"],),
        ).fetchone()[0]
        orphans = connection.execute(
            """SELECT COUNT(*) FROM chunks c
            LEFT JOIN chunk_blocks cb ON cb.chunk_id=c.id
            WHERE c.document_id=? AND cb.block_id IS NULL""",
            (imported["document_id"],),
        ).fetchone()[0]
        versions = connection.execute(
            "SELECT DISTINCT chunker_version FROM chunks WHERE document_id=?",
            (imported["document_id"],),
        ).fetchall()
    assert blocks >= 2
    assert orphans == 0
    assert [row[0] for row in versions] == ["typed-v2"]


def test_text_metrics_report_recall_and_precision_separately() -> None:
    metrics = text_metrics("abc", "abx")
    assert metrics == {
        "text_recall": 0.666667,
        "text_precision": 0.666667,
        "text_f1": 0.666667,
    }


def test_reindex_outdated_preserves_source_and_document_identity(
    knowledge_system: KnowledgeSystem,
    source_root: Path,
) -> None:
    note = source_root / "architecture.md"
    note.write_text("# 系统架构\n\n混合解析保留来源定位。", encoding="utf-8")
    imported = knowledge_system.ingestion.import_file(
        note,
        domain="work",
        privacy="private",
    )
    source_id = imported["source_id"]
    document_id = imported["document_id"]
    with knowledge_system.database.connect() as connection:
        connection.execute(
            "UPDATE documents SET parser_version = '1' WHERE id = ?",
            (document_id,),
        )
        connection.commit()

    result = knowledge_system.ingestion.reindex_outdated()
    document = knowledge_system.repository.read_document(document_id)
    search = knowledge_system.repository.search("混合解析保留来源定位", domain="work")

    assert result["status"] == "completed"
    assert result["reindexed"] == 1
    assert document is not None
    assert document["source_id"] == source_id
    assert document["title"] == "architecture"
    assert document["metadata"]["extraction"]["pipeline_version"] == "hybrid-v2"
    assert search[0]["document_id"] == document_id


def test_forced_targeted_reindex_only_rebuilds_requested_source_type(
    knowledge_system: KnowledgeSystem,
    source_root: Path,
) -> None:
    pdf_path = source_root / "target.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "目标 PDF 重建")
    document.save(pdf_path)
    document.close()
    note_path = source_root / "note.md"
    note_path.write_text("不应被 PDF 定向重建", encoding="utf-8")
    knowledge_system.ingestion.import_file(pdf_path, domain="work", privacy="private")
    note = knowledge_system.ingestion.import_file(note_path, domain="work", privacy="private")
    with knowledge_system.database.connect() as connection:
        before = connection.execute(
            "SELECT created_at FROM documents WHERE id=?", (note["document_id"],)
        ).fetchone()[0]

    result = knowledge_system.ingestion.reindex_outdated(
        source_types={"pdf"}, force=True
    )

    with knowledge_system.database.connect() as connection:
        after = connection.execute(
            "SELECT created_at FROM documents WHERE id=?", (note["document_id"],)
        ).fetchone()[0]
    assert result["status"] == "completed"
    assert result["candidates"] == result["reindexed"] == 1
    assert result["source_types"] == ["pdf"]
    assert result["forced"] is True
    assert before == after
