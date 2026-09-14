import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from pkas.system import KnowledgeSystem
from pkas.weflow import WeFlowFormatError

CHATLAB_FIXTURE = Path(__file__).parent / "fixtures" / "weflow_chatlab_private.json"


def copy_chatlab_fixture(source_root: Path) -> Path:
    target = source_root / "weflow-customer.json"
    target.write_bytes(CHATLAB_FIXTURE.read_bytes())
    return target


def create_xlsx_export_fixture(source_root: Path) -> tuple[Path, Path]:
    xlsx = source_root / "测试客户张经理-聊天记录.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["微信聊天记录"])
    sheet.append(["导出工具", "WeFlow", "导出版本", "test"])
    sheet.append(["昵称", "测试客户张经理", "微信ID", "wxid_customer_demo"])
    sheet.append([])
    sheet.append(
        [
            "序号",
            "时间",
            "发送者昵称",
            "发送者微信ID",
            "发送者备注",
            "发送者身份",
            "消息类型",
            "内容",
        ]
    )
    sheet.append(
        [
            1,
            "2025-02-05 09:20:00",
            "张经理",
            "wxid_customer_demo",
            "客户-张经理",
            "客户",
            "文本消息",
            "门店系统希望本周五前看到报价单。",
        ]
    )
    sheet.append(
        [
            2,
            "2025-02-05 09:21:00",
            "我",
            "wxid_owner",
            "",
            "我",
            "文本消息",
            "收到，我会先核对接口范围再给您正式报价。",
        ]
    )
    sheet.append(
        [
            3,
            "2025-02-05 09:22:00",
            "张经理",
            "wxid_customer_demo",
            "客户-张经理",
            "客户",
            "文本消息",
            "还需要包含三家门店的数据看板。",
        ]
    )
    workbook.save(xlsx)
    records = source_root / "weflow-export-records.json"
    records.write_text(
        json.dumps(
            {
                "wxid_customer_demo": [
                    {
                        "exportTime": 1738713900000,
                        "format": "xlsx",
                        "messageCount": 3,
                        "outputPath": str(xlsx.resolve()),
                    }
                ]
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    return records, xlsx


def clone_xlsx_for_session(template: Path, session_id: str) -> Path:
    target = template.with_name(f"{session_id}.xlsx")
    workbook = load_workbook(template)
    sheet = workbook.active
    assert sheet is not None
    sheet.cell(row=3, column=4, value=session_id)
    workbook.save(target)
    return target


def test_weflow_xlsx_export_discovery_import_deduplication_and_context(
    knowledge_system: KnowledgeSystem,
    source_root: Path,
) -> None:
    records, _ = create_xlsx_export_fixture(source_root)
    catalog = knowledge_system.weflow.discover_exports(records_path=str(records))
    assert catalog["api_required"] is False
    assert catalog["key_accessed"] is False
    assert catalog["total_sessions"] == 1
    assert catalog["existing_sessions"] == 1
    assert catalog["existing_record_count"] == 1

    inspection = knowledge_system.weflow.inspect_export_selection(
        records_path=str(records),
        session_ids=["wxid_customer_demo"],
    )
    assert inspection["selected_sessions"] == 1
    assert inspection["total_messages"] == 3
    assert inspection["items"][0]["display_name"] == "测试客户张经理"
    assert inspection["items"][0]["header_row"] == 5

    imported = knowledge_system.customer_workflows.import_weflow_exports(
        records_path=str(records),
        session_ids=["wxid_customer_demo"],
        inspection_token=inspection["inspection_token"],
        privacy="restricted",
    )
    assert imported["status"] == "completed"
    assert imported["result"]["imported"] == 3
    snapshot = Path(imported["result"]["snapshot_paths"][0])
    assert snapshot.is_file()
    assert snapshot.suffix == ".xlsx"

    customers = knowledge_system.customers.list_customers()
    assert len(customers) == 1
    customer = customers[0]
    assert customer["display_name"] == "测试客户张经理"
    assert customer["message_count"] == 3
    assert knowledge_system.customers.timeline(customer["id"]) == []
    timeline = knowledge_system.customers.timeline(
        customer["id"],
        include_restricted=True,
    )
    assert len(timeline) == 3
    assert timeline[0]["content"] == "还需要包含三家门店的数据看板。"
    assert timeline[0]["local_id"] == "xlsx-row:8"
    assert timeline[1]["is_self"] == 1
    assert timeline[0]["source_uri"].endswith("测试客户张经理-聊天记录.xlsx")

    results = knowledge_system.customers.search_messages(
        "三家门店 数据看板",
        customer_id=customer["id"],
        include_restricted=True,
    )
    assert len(results) == 1
    assert results[0]["context_count"] == 3
    assert results[0]["context_window"] == "adjacent-3-v1"
    assert [item["is_match"] for item in results[0]["context_messages"]] == [
        False,
        False,
        True,
    ]
    assert results[0]["context_messages"][0]["content"] == "门店系统希望本周五前看到报价单。"
    approved = knowledge_system.customers.update_customer(
        customer["id"],
        review_status="approved",
    )
    assert approved is not None
    context = knowledge_system.customer_service.prepare_reply_context(
        customer_id=customer["id"],
        task="回复客户关于门店数据看板和报价的问题",
        include_restricted=True,
    )
    assert context is not None
    assert len(context["recent_messages"]) == 3

    duplicate = knowledge_system.customer_workflows.import_weflow_exports(
        records_path=str(records),
        session_ids=["wxid_customer_demo"],
        inspection_token=inspection["inspection_token"],
        privacy="restricted",
    )
    assert duplicate["result"]["imported"] == 0
    assert duplicate["result"]["duplicates"] == 3

    with knowledge_system.database.connect() as connection:
        connector = connection.execute(
            "SELECT config_json FROM connectors WHERE connector_type = 'weflow-xlsx-export'"
        ).fetchone()
    assert '"api_required":false' in connector["config_json"]
    assert '"key_accessed":false' in connector["config_json"]


def test_weflow_xlsx_inspection_token_detects_file_change(
    knowledge_system: KnowledgeSystem,
    source_root: Path,
) -> None:
    records, xlsx = create_xlsx_export_fixture(source_root)
    inspection = knowledge_system.weflow.inspect_export_selection(
        records_path=str(records),
        session_ids=["wxid_customer_demo"],
    )
    workbook = load_workbook(xlsx)
    sheet = workbook.active
    assert sheet is not None
    sheet.append(
        [
            4,
            "2025-02-05 09:23:00",
            "张经理",
            "wxid_customer_demo",
            "客户-张经理",
            "客户",
            "文本消息",
            "检查后新增的消息",
        ]
    )
    workbook.save(xlsx)

    result = knowledge_system.customer_workflows.import_weflow_exports(
        records_path=str(records),
        session_ids=["wxid_customer_demo"],
        inspection_token=inspection["inspection_token"],
        privacy="restricted",
    )
    assert result["status"] == "failed"
    assert "发生了变化" in result["error"]["message"]


def test_weflow_xlsx_supports_more_than_one_hundred_sessions(
    knowledge_system: KnowledgeSystem,
    source_root: Path,
) -> None:
    records, xlsx = create_xlsx_export_fixture(source_root)
    session_ids = [f"wxid_customer_{index:03d}" for index in range(101)]
    export_paths = {
        session_id: clone_xlsx_for_session(xlsx, session_id) for session_id in session_ids
    }
    records.write_text(
        json.dumps(
            {
                session_id: [
                    {
                        "exportTime": 1738713900000,
                        "format": "xlsx",
                        "messageCount": 3,
                        "outputPath": str(export_paths[session_id].resolve()),
                    }
                ]
                for session_id in session_ids
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    inspection = knowledge_system.weflow.inspect_export_selection(
        records_path=str(records),
        session_ids=session_ids,
    )

    assert inspection["selected_sessions"] == 101
    assert inspection["total_messages"] == 303


def test_weflow_batch_import_continues_after_one_session_failure(
    knowledge_system: KnowledgeSystem,
    source_root: Path,
    monkeypatch,
) -> None:
    records, xlsx = create_xlsx_export_fixture(source_root)
    session_ids = ["wxid_customer_demo", "wxid_customer_failure"]
    failure_xlsx = clone_xlsx_for_session(xlsx, "wxid_customer_failure")
    records.write_text(
        json.dumps(
            {
                session_id: [
                    {
                        "exportTime": 1738713900000,
                        "format": "xlsx",
                        "messageCount": 3,
                        "outputPath": str(
                            xlsx.resolve()
                            if session_id == "wxid_customer_demo"
                            else failure_xlsx.resolve()
                        ),
                    }
                ]
                for session_id in session_ids
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    inspection = knowledge_system.weflow.inspect_export_selection(
        records_path=str(records),
        session_ids=session_ids,
    )
    original_import = knowledge_system.weflow._import_xlsx_export

    def import_with_one_failure(*, connector_id, item, privacy):
        if item["session_id"] == "wxid_customer_failure":
            raise OSError("simulated session failure")
        return original_import(connector_id=connector_id, item=item, privacy=privacy)

    monkeypatch.setattr(knowledge_system.weflow, "_import_xlsx_export", import_with_one_failure)
    result = knowledge_system.customer_workflows.import_weflow_exports(
        records_path=str(records),
        session_ids=session_ids,
        inspection_token=inspection["inspection_token"],
        privacy="restricted",
    )

    assert result["status"] == "warning"
    assert result["result"]["imported"] == 3
    assert result["result"]["failed_sessions"] == 1
    assert result["result"]["session_errors"][0]["session_id"] == "wxid_customer_failure"


def test_weflow_duplicate_record_alias_uses_xlsx_metadata_canonical_session(
    knowledge_system: KnowledgeSystem,
    source_root: Path,
) -> None:
    records, xlsx = create_xlsx_export_fixture(source_root)
    record = {
        "exportTime": 1738713900000,
        "format": "xlsx",
        "messageCount": 3,
        "outputPath": str(xlsx.resolve()),
    }
    records.write_text(
        json.dumps(
            {
                "wxid_customer_demo": [record],
                "wxid_stale_alias": [record],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    inspection = knowledge_system.weflow.inspect_export_selection(
        records_path=str(records),
        session_ids=["wxid_customer_demo", "wxid_stale_alias"],
    )

    assert inspection["requested_sessions"] == 2
    assert inspection["selected_sessions"] == 1
    assert inspection["skipped_alias_sessions"] == 1
    assert inspection["canonical_session_ids"] == ["wxid_customer_demo"]

    imported = knowledge_system.weflow.import_export_selection(
        records_path=str(records),
        session_ids=["wxid_customer_demo", "wxid_stale_alias"],
        inspection_token=inspection["inspection_token"],
    )
    assert imported["imported"] == 3
    assert imported["skipped_alias_sessions"] == 1
    assert len(knowledge_system.customers.list_customers()) == 1


def test_weflow_metadata_mismatch_without_canonical_record_is_rejected(
    knowledge_system: KnowledgeSystem,
    source_root: Path,
) -> None:
    records, xlsx = create_xlsx_export_fixture(source_root)
    records.write_text(
        json.dumps(
            {
                "wxid_wrong_alias": [
                    {
                        "exportTime": 1738713900000,
                        "format": "xlsx",
                        "messageCount": 3,
                        "outputPath": str(xlsx.resolve()),
                    }
                ]
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    with pytest.raises(WeFlowFormatError, match="会话 ID 与 XLSX 元数据不一致"):
        knowledge_system.weflow.inspect_export_selection(
            records_path=str(records),
            session_ids=["wxid_wrong_alias"],
        )


def test_weflow_batch_import_reports_when_every_session_fails(
    knowledge_system: KnowledgeSystem,
    source_root: Path,
    monkeypatch,
) -> None:
    records, _ = create_xlsx_export_fixture(source_root)
    inspection = knowledge_system.weflow.inspect_export_selection(
        records_path=str(records),
        session_ids=["wxid_customer_demo"],
    )

    def fail_import(*, connector_id, item, privacy):
        raise OSError("simulated complete failure")

    monkeypatch.setattr(knowledge_system.weflow, "_import_xlsx_export", fail_import)
    result = knowledge_system.customer_workflows.import_weflow_exports(
        records_path=str(records),
        session_ids=["wxid_customer_demo"],
        inspection_token=inspection["inspection_token"],
        privacy="restricted",
    )

    assert result["status"] == "failed"
    assert result["result"]["failed_sessions"] == 1
    assert result["error"]["type"] == "WeFlowBatchImportError"
    assert result["error"]["safe_retry"]


def test_offline_chatlab_import_remains_supported(
    knowledge_system: KnowledgeSystem,
    source_root: Path,
) -> None:
    chatlab = copy_chatlab_fixture(source_root)
    inspection = knowledge_system.weflow.inspect_chatlab_file(str(chatlab))
    assert inspection["session_id"] == "wxid_customer_demo"
    assert inspection["message_count"] == 3

    imported = knowledge_system.customer_workflows.import_chatlab(
        path=str(chatlab),
        inspection_token=inspection["inspection_token"],
        session_id=None,
        privacy="restricted",
    )
    assert imported["status"] == "completed"
    assert imported["result"]["imported"] == 3


def test_chatlab_inspection_token_detects_file_change(
    knowledge_system: KnowledgeSystem,
    source_root: Path,
) -> None:
    chatlab = copy_chatlab_fixture(source_root)
    inspection = knowledge_system.weflow.inspect_chatlab_file(str(chatlab))
    payload = json.loads(chatlab.read_text(encoding="utf-8"))
    payload["messages"].append(
        {
            "sender": "wxid_customer_demo",
            "timestamp": 1738714000,
            "type": 0,
            "content": "检查后新增的消息",
            "platformMessageId": "wf-demo-004",
        }
    )
    chatlab.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")

    result = knowledge_system.customer_workflows.import_chatlab(
        path=str(chatlab),
        inspection_token=inspection["inspection_token"],
        session_id=None,
        privacy="restricted",
    )
    assert result["status"] == "failed"
    assert "发生了变化" in result["error"]["message"]
