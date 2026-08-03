import hashlib
import json
import shutil
import time
from collections.abc import Iterator
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from pkas.config import Settings, get_settings
from pkas.customer_repository import CustomerRepository
from pkas.ingest import ImportBoundaryError, IngestionService

XLSX_HEADERS = (
    "序号",
    "时间",
    "发送者昵称",
    "发送者微信ID",
    "发送者备注",
    "发送者身份",
    "消息类型",
    "内容",
)
XLSX_REQUIRED_HEADERS = {
    "时间",
    "发送者昵称",
    "发送者微信ID",
    "发送者身份",
    "消息类型",
    "内容",
}
XLSX_METADATA_KEYS = {
    "昵称",
    "微信ID",
    "群聊名称",
    "群名称",
    "聊天对象",
    "导出时间",
    "导出工具",
    "导出版本",
}


class WeFlowError(RuntimeError):
    pass


class WeFlowFormatError(ValueError):
    pass


def validate_chatlab(payload: dict[str, Any]) -> None:
    chatlab = payload.get("chatlab")
    meta = payload.get("meta")
    messages = payload.get("messages")
    if (
        not isinstance(chatlab, dict)
        or not isinstance(meta, dict)
        or not isinstance(messages, list)
    ):
        raise WeFlowFormatError("文件不是有效的 ChatLab 会话结构。")
    if str(chatlab.get("generator", "")).lower() != "weflow":
        raise WeFlowFormatError("ChatLab 文件不是由 WeFlow 生成的。")
    if str(meta.get("platform", "wechat")).lower() != "wechat":
        raise WeFlowFormatError("ChatLab 文件不是微信会话。")


def detect_session_id(payload: dict[str, Any], fallback: str | None = None) -> str:
    meta = payload.get("meta", {})
    for key in ("groupId", "contactId", "sessionId", "talker", "id"):
        value = meta.get(key) if isinstance(meta, dict) else None
        if value:
            return str(value)
    owner_id = str(meta.get("ownerId", "")) if isinstance(meta, dict) else ""
    senders = {
        str(message.get("sender"))
        for message in payload.get("messages", [])
        if isinstance(message, dict)
        and message.get("sender")
        and str(message.get("sender")) != owner_id
    }
    if len(senders) == 1:
        return senders.pop()
    if fallback:
        return fallback
    raise WeFlowFormatError("无法确定私聊会话 ID，请明确填写微信 wxid。")


def _file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _timestamp(value: Any) -> int:
    if isinstance(value, datetime):
        return int(time.mktime(value.timetuple()))
    if isinstance(value, (int, float)):
        numeric = int(value)
        return numeric // 1000 if numeric > 10_000_000_000 else numeric
    text = str(value or "").strip()
    if not text:
        return 0
    if text.isdigit():
        numeric = int(text)
        return numeric // 1000 if numeric > 10_000_000_000 else numeric
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
        if parsed.tzinfo:
            return int(parsed.timestamp())
        return int(time.mktime(parsed.timetuple()))
    except ValueError:
        return 0


def _record_timestamp(value: Any) -> int:
    try:
        numeric = int(value)
    except (TypeError, ValueError):
        return 0
    return numeric // 1000 if numeric > 10_000_000_000 else numeric


def _row_value(row: tuple[Any, ...], indexes: dict[str, int], name: str) -> Any:
    position = indexes.get(name)
    if position is None or position >= len(row):
        return None
    return row[position]


class WeFlowService:
    def __init__(
        self,
        *,
        settings: Settings | None = None,
        customers: CustomerRepository,
        ingestion: IngestionService,
    ) -> None:
        self.settings = settings or get_settings()
        self.customers = customers
        self.ingestion = ingestion

    def discover_exports(
        self,
        *,
        records_path: str | None = None,
        keyword: str = "",
        limit: int = 500,
    ) -> dict[str, Any]:
        catalog_path, payload = self._load_export_records(records_path)
        items: list[dict[str, Any]] = []
        record_count = 0
        existing_record_count = 0
        existing_session_count = 0
        for session_id, raw_records in payload.items():
            records = [item for item in raw_records if isinstance(item, dict)]
            record_count += len(records)
            existing_record_count += sum(
                1
                for record in records
                if self._record_source(record, require_exists=False).is_file()
            )
            selected = self._latest_existing_record(records)
            if not selected:
                continue
            existing_session_count += 1
            source = self._record_source(selected)
            display_name = source.stem
            search_text = f"{session_id} {display_name}".lower()
            if keyword.strip() and keyword.strip().lower() not in search_text:
                continue
            items.append(
                {
                    "session_id": str(session_id),
                    "display_name": display_name,
                    "conversation_type": (
                        "group" if str(session_id).endswith("@chatroom") else "private"
                    ),
                    "format": "xlsx",
                    "export_time": _record_timestamp(selected.get("exportTime")),
                    "message_count": int(selected.get("messageCount") or 0),
                    "output_path": str(source),
                    "byte_size": source.stat().st_size,
                    "existing_export_count": sum(
                        1 for record in records if self._record_source(record, False).is_file()
                    ),
                }
            )
        items.sort(key=lambda item: (item["export_time"], item["session_id"]), reverse=True)
        return {
            "records_path": str(catalog_path),
            "total_sessions": len(payload),
            "existing_sessions": existing_session_count,
            "missing_sessions": len(payload) - existing_session_count,
            "matched_sessions": len(items),
            "record_count": record_count,
            "existing_record_count": existing_record_count,
            "items": items[: max(1, min(limit, 1000))],
            "api_required": False,
            "key_accessed": False,
        }

    def inspect_export_selection(
        self,
        *,
        records_path: str | None,
        session_ids: list[str],
    ) -> dict[str, Any]:
        catalog_path, payload = self._load_export_records(records_path)
        selected = self._select_exports(payload, session_ids)
        records_hash = _file_hash(catalog_path)
        inspected: list[dict[str, Any]] = []
        token_items: list[dict[str, str]] = []
        for session_id, record, source in selected:
            layout = self._inspect_xlsx_layout(source)
            content_hash = _file_hash(source)
            display_name = layout["display_name"] or source.stem
            inspected.append(
                {
                    "session_id": session_id,
                    "display_name": display_name,
                    "conversation_type": (
                        "group" if session_id.endswith("@chatroom") else "private"
                    ),
                    "output_path": str(source),
                    "byte_size": source.stat().st_size,
                    "content_hash": content_hash,
                    "message_count": int(record.get("messageCount") or 0),
                    "export_time": _record_timestamp(record.get("exportTime")),
                    "header_row": layout["header_row"],
                    "sheet_count": layout["sheet_count"],
                }
            )
            token_items.append(
                {
                    "session_id": session_id,
                    "path": str(source),
                    "content_hash": content_hash,
                }
            )
        inspection_token = hashlib.sha256(
            json.dumps(
                {
                    "records_path": str(catalog_path),
                    "records_hash": records_hash,
                    "items": token_items,
                },
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            ).encode("utf-8")
        ).hexdigest()
        return {
            "records_path": str(catalog_path),
            "selected_sessions": len(inspected),
            "total_messages": sum(item["message_count"] for item in inspected),
            "total_bytes": sum(item["byte_size"] for item in inspected),
            "items": inspected,
            "inspection_token": inspection_token,
            "privacy": "restricted",
            "api_required": False,
            "key_accessed": False,
        }

    def import_export_selection(
        self,
        *,
        records_path: str | None,
        session_ids: list[str],
        inspection_token: str,
        privacy: str = "restricted",
    ) -> dict[str, Any]:
        inspection = self.inspect_export_selection(
            records_path=records_path,
            session_ids=session_ids,
        )
        if inspection["inspection_token"] != inspection_token:
            raise ImportBoundaryError("WeFlow 导出记录或 XLSX 文件在检查后发生了变化，请重新检查。")
        connector_id = self.customers.upsert_connector(
            connector_type="weflow-xlsx-export",
            name="WeFlow XLSX 导出记录",
            base_url="offline://weflow-export-records",
            status="ready",
            config={"transport": "xlsx-export", "api_required": False, "key_accessed": False},
        )
        results: list[dict[str, Any]] = []
        session_errors: list[dict[str, Any]] = []
        for item in inspection["items"]:
            try:
                results.append(
                    self._import_xlsx_export(
                        connector_id=connector_id,
                        item=item,
                        privacy=privacy,
                    )
                )
            except Exception as exc:
                session_errors.append(
                    {
                        "session_id": item["session_id"],
                        "display_name": item["display_name"],
                        "error_type": type(exc).__name__,
                        "message": str(exc),
                    }
                )
        return {
            "connector_id": connector_id,
            "sessions": results,
            "failed_sessions": len(session_errors),
            "session_errors": session_errors,
            "imported": sum(item["imported"] for item in results),
            "duplicates": sum(item["duplicates"] for item in results),
            "snapshot_paths": [item["snapshot_path"] for item in results],
            "api_required": False,
            "key_accessed": False,
        }

    def _import_xlsx_export(
        self,
        *,
        connector_id: str,
        item: dict[str, Any],
        privacy: str,
    ) -> dict[str, Any]:
        source = Path(item["output_path"])
        snapshot = self._store_file_snapshot(
            source,
            expected_hash=item["content_hash"],
            metadata={
                "session_id": item["session_id"],
                "source_format": "weflow-xlsx",
                "export_time": item["export_time"],
                "declared_message_count": item["message_count"],
            },
        )
        snapshot_source = Path(snapshot["vault_path"])
        layout = self._inspect_xlsx_layout(snapshot_source)
        layout["session_id"] = item["session_id"]
        imported = 0
        duplicates = 0
        pages = 0
        watermark = 0
        last_result: dict[str, Any] | None = None
        for messages, members in self._iter_xlsx_pages(snapshot_source, layout):
            payload = {
                "chatlab": {"version": "weflow-xlsx-1", "generator": "WeFlow"},
                "meta": {
                    "name": layout["display_name"] or item["display_name"],
                    "platform": "wechat",
                    "type": item["conversation_type"],
                    "sessionId": item["session_id"],
                },
                "members": members,
                "messages": messages,
            }
            last_result = self.customers.ingest_chatlab_page(
                payload=payload,
                session_id=item["session_id"],
                connector_id=connector_id,
                snapshot=snapshot,
                privacy=privacy,
                contact={
                    "displayName": layout["display_name"] or item["display_name"],
                    "nickname": layout["display_name"] or item["display_name"],
                },
            )
            imported += last_result["imported"]
            duplicates += last_result["duplicates"]
            pages += 1
            watermark = max(watermark, *(message["timestamp"] for message in messages))
        if last_result is None:
            raise WeFlowFormatError("WeFlow XLSX 中没有可导入的聊天消息。")
        self.customers.update_sync_cursor(
            item["session_id"],
            since=watermark,
            offset=0,
            watermark=watermark,
        )
        return {
            "session_id": item["session_id"],
            "display_name": layout["display_name"] or item["display_name"],
            "customer_id": last_result["customer_id"],
            "conversation_id": last_result["conversation_id"],
            "imported": imported,
            "duplicates": duplicates,
            "message_count": last_result["message_count"],
            "pages": pages,
            "snapshot_path": snapshot["vault_path"],
        }

    def _iter_xlsx_pages(
        self,
        source: Path,
        layout: dict[str, Any],
        page_size: int = 2000,
    ) -> Iterator[tuple[list[dict[str, Any]], list[dict[str, Any]]]]:
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            sheet = workbook.worksheets[0]
            indexes = layout["indexes"]
            messages: list[dict[str, Any]] = []
            members: dict[str, dict[str, Any]] = {}
            for row_number, row in enumerate(
                sheet.iter_rows(min_row=layout["header_row"] + 1, values_only=True),
                start=layout["header_row"] + 1,
            ):
                sender_id = str(_row_value(row, indexes, "发送者微信ID") or "").strip()
                sender_name = str(_row_value(row, indexes, "发送者昵称") or sender_id).strip()
                sender_remark = str(_row_value(row, indexes, "发送者备注") or "").strip()
                identity = str(_row_value(row, indexes, "发送者身份") or "").strip()
                message_type = str(_row_value(row, indexes, "消息类型") or "unknown").strip()
                content = str(_row_value(row, indexes, "内容") or "").strip()
                sent_at = _timestamp(_row_value(row, indexes, "时间"))
                if not any((sender_id, sender_name, content, sent_at)):
                    continue
                dedup_key = json.dumps(
                    {
                        "session_id": layout["session_id"],
                        "sender": sender_id,
                        "timestamp": sent_at,
                        "type": message_type,
                        "content": content,
                    },
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                )
                messages.append(
                    {
                        "sender": sender_id,
                        "accountName": sender_remark or sender_name,
                        "timestamp": sent_at,
                        "type": message_type,
                        "content": content,
                        "isSend": identity == "我",
                        "localId": f"xlsx-row:{row_number}",
                        "dedupKey": dedup_key,
                    }
                )
                if sender_id:
                    members[sender_id] = {
                        "platformId": sender_id,
                        "accountName": sender_remark or sender_name,
                        "isOwner": identity == "我",
                    }
                if len(messages) >= page_size:
                    yield messages, list(members.values())
                    messages = []
                    members = {}
            if messages:
                yield messages, list(members.values())
        finally:
            workbook.close()

    def _inspect_xlsx_layout(self, source: Path) -> dict[str, Any]:
        if source.suffix.lower() != ".xlsx" or not source.is_file():
            raise WeFlowFormatError("WeFlow 导出文件必须是现存的 XLSX 文件。")
        if source.stat().st_size > self.settings.max_source_bytes:
            raise WeFlowFormatError(
                f"WeFlow XLSX 超过单文件上限 {self.settings.max_source_bytes} 字节。"
            )
        try:
            workbook = load_workbook(source, read_only=True, data_only=True)
        except (OSError, ValueError) as exc:
            raise WeFlowFormatError(f"无法打开 WeFlow XLSX：{exc}") from exc
        try:
            if not workbook.worksheets:
                raise WeFlowFormatError("WeFlow XLSX 没有工作表。")
            sheet = workbook.worksheets[0]
            metadata: dict[str, str] = {}
            headers: list[str] | None = None
            header_row = 0
            for row_number, row in enumerate(
                sheet.iter_rows(min_row=1, max_row=40, values_only=True),
                start=1,
            ):
                values = [str(value).strip() if value is not None else "" for value in row]
                if XLSX_REQUIRED_HEADERS.issubset(set(values)):
                    headers = values
                    header_row = row_number
                    break
                for index, label in enumerate(values[:-1]):
                    if label in XLSX_METADATA_KEYS and values[index + 1]:
                        metadata[label] = values[index + 1]
            if headers is None:
                raise WeFlowFormatError("XLSX 中未找到 WeFlow 标准聊天表头。")
            indexes = {name: headers.index(name) for name in XLSX_HEADERS if name in headers}
            display_name = (
                metadata.get("昵称")
                or metadata.get("群聊名称")
                or metadata.get("群名称")
                or metadata.get("聊天对象")
                or ""
            )
            return {
                "header_row": header_row,
                "indexes": indexes,
                "display_name": display_name,
                "metadata_session_id": metadata.get("微信ID") or "",
                "sheet_count": len(workbook.worksheets),
                "session_id": metadata.get("微信ID") or "",
            }
        finally:
            workbook.close()

    def _load_export_records(
        self,
        records_path: str | None,
    ) -> tuple[Path, dict[str, list[Any]]]:
        raw_path = records_path or str(self.settings.weflow_export_records_path)
        path = Path(raw_path).expanduser()
        if not path.is_absolute():
            raise WeFlowFormatError("WeFlow 导出记录必须使用绝对路径。")
        path = path.resolve()
        if not path.is_file():
            raise WeFlowFormatError(f"WeFlow 导出记录不存在：{path}")
        if path.stat().st_size > self.settings.max_source_bytes:
            raise WeFlowFormatError("WeFlow 导出记录文件超过读取上限。")
        try:
            payload = json.loads(path.read_text(encoding="utf-8-sig"))
        except (OSError, json.JSONDecodeError) as exc:
            raise WeFlowFormatError(f"无法读取 WeFlow 导出记录：{exc}") from exc
        if not isinstance(payload, dict) or not all(
            isinstance(key, str) and isinstance(value, list) for key, value in payload.items()
        ):
            raise WeFlowFormatError("WeFlow 导出记录结构无效。")
        return path, payload

    @staticmethod
    def _record_source(record: dict[str, Any], require_exists: bool = True) -> Path:
        raw_path = str(record.get("outputPath") or "")
        path = Path(raw_path).expanduser()
        if not path.is_absolute() or path.suffix.lower() != ".xlsx":
            return Path("__invalid_weflow_export__")
        path = path.resolve()
        if require_exists and not path.is_file():
            raise WeFlowFormatError(f"WeFlow 导出文件不存在：{path}")
        return path

    def _latest_existing_record(self, records: list[dict[str, Any]]) -> dict[str, Any] | None:
        ordered = sorted(
            records,
            key=lambda item: _record_timestamp(item.get("exportTime")),
            reverse=True,
        )
        for record in ordered:
            if self._record_source(record, require_exists=False).is_file():
                return record
        return None

    def _select_exports(
        self,
        payload: dict[str, list[Any]],
        session_ids: list[str],
    ) -> list[tuple[str, dict[str, Any], Path]]:
        unique_ids = list(dict.fromkeys(session_ids))
        if not unique_ids:
            raise WeFlowFormatError("至少选择一个 WeFlow 导出会话。")
        if len(unique_ids) > 1000:
            raise WeFlowFormatError("单次最多导入 1000 个 WeFlow 导出会话。")
        selected: list[tuple[str, dict[str, Any], Path]] = []
        for session_id in unique_ids:
            raw_records = payload.get(session_id)
            if not isinstance(raw_records, list):
                raise WeFlowFormatError(f"导出记录中不存在所选会话：{session_id}")
            record = self._latest_existing_record(
                [item for item in raw_records if isinstance(item, dict)]
            )
            if not record:
                raise WeFlowFormatError(f"所选会话没有现存的 XLSX 导出：{session_id}")
            selected.append((session_id, record, self._record_source(record)))
        return selected

    def inspect_chatlab_file(
        self,
        path: str,
        *,
        session_id: str | None = None,
    ) -> dict[str, Any]:
        inspection = self.ingestion.inspect_path(path, recursive=False)
        target = Path(inspection["path"])
        if not target.is_file() or target.suffix.lower() != ".json":
            raise WeFlowFormatError("请选择一个 WeFlow ChatLab JSON 文件。")
        if inspection["total_bytes"] > self.settings.max_source_bytes:
            raise WeFlowFormatError(
                f"ChatLab 文件超过单文件上限 {self.settings.max_source_bytes} 字节。"
            )
        payload = self._load_chatlab(target)
        detected_session_id = detect_session_id(payload, session_id)
        return {
            **inspection,
            "session_id": detected_session_id,
            "name": payload.get("meta", {}).get("name") or detected_session_id,
            "conversation_type": payload.get("meta", {}).get("type") or "private",
            "message_count": len(payload.get("messages", [])),
            "chatlab_version": payload.get("chatlab", {}).get("version"),
            "generator": payload.get("chatlab", {}).get("generator"),
        }

    def import_chatlab_file(
        self,
        *,
        path: str,
        inspection_token: str,
        session_id: str | None = None,
        privacy: str = "restricted",
    ) -> dict[str, Any]:
        inspection = self.inspect_chatlab_file(path, session_id=session_id)
        if inspection["inspection_token"] != inspection_token:
            raise ImportBoundaryError("ChatLab 文件在检查后发生了变化，请重新检查后再确认导入。")
        target = Path(inspection["path"])
        payload = self._load_chatlab(target)
        detected_session_id = detect_session_id(payload, session_id)
        connector_id = self.customers.upsert_connector(
            connector_type="weflow-chatlab",
            name="WeFlow ChatLab 离线导入",
            base_url="offline://weflow-chatlab",
            status="ready",
            config={"transport": "file", "api_required": False, "key_accessed": False},
        )
        snapshot = self._store_json_snapshot(
            payload,
            source_uri=str(target.resolve()),
            metadata={"session_id": detected_session_id, "offline_import": True},
        )
        result = self.customers.ingest_chatlab_page(
            payload=payload,
            session_id=detected_session_id,
            connector_id=connector_id,
            snapshot=snapshot,
            privacy=privacy,
        )
        watermark = max(
            (int(message.get("timestamp") or 0) for message in payload.get("messages", [])),
            default=0,
        )
        self.customers.update_sync_cursor(
            detected_session_id,
            since=watermark,
            offset=0,
            watermark=watermark,
        )
        return {
            **result,
            "session_id": detected_session_id,
            "snapshot_path": snapshot["vault_path"],
            "api_required": False,
            "key_accessed": False,
        }

    @staticmethod
    def _load_chatlab(path: Path) -> dict[str, Any]:
        try:
            payload = json.loads(path.read_text(encoding="utf-8-sig"))
        except (OSError, json.JSONDecodeError) as exc:
            raise WeFlowFormatError(f"无法读取 WeFlow ChatLab JSON：{exc}") from exc
        if not isinstance(payload, dict):
            raise WeFlowFormatError("ChatLab 文件顶层必须是 JSON 对象。")
        validate_chatlab(payload)
        return payload

    def _store_file_snapshot(
        self,
        source: Path,
        *,
        expected_hash: str,
        metadata: dict[str, Any],
    ) -> dict[str, Any]:
        content_hash = _file_hash(source)
        if content_hash != expected_hash:
            raise ImportBoundaryError("WeFlow XLSX 在检查后发生了变化，请重新检查。")
        vault_dir = (
            self.settings.data_root / "raw" / "weflow-xlsx" / "sha256" / content_hash[:2]
        )
        vault_dir.mkdir(parents=True, exist_ok=True)
        vault_path = vault_dir / f"{content_hash}.xlsx"
        if not vault_path.exists():
            shutil.copy2(source, vault_path)
        if _file_hash(vault_path) != content_hash:
            raise OSError("WeFlow XLSX 快照写入后的哈希校验失败。")
        return {
            "content_hash": content_hash,
            "vault_path": str(vault_path.resolve()),
            "byte_size": source.stat().st_size,
            "source_uri": str(source.resolve()),
            "metadata": metadata,
        }

    def _store_json_snapshot(
        self,
        payload: dict[str, Any],
        *,
        source_uri: str,
        metadata: dict[str, Any],
    ) -> dict[str, Any]:
        encoded = json.dumps(
            payload,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
        content_hash = hashlib.sha256(encoded).hexdigest()
        vault_dir = self.settings.data_root / "raw" / "weflow" / "sha256" / content_hash[:2]
        vault_dir.mkdir(parents=True, exist_ok=True)
        vault_path = vault_dir / f"{content_hash}.json"
        if not vault_path.exists():
            vault_path.write_bytes(encoded)
        if hashlib.sha256(vault_path.read_bytes()).hexdigest() != content_hash:
            raise OSError("WeFlow 快照写入后的哈希校验失败。")
        return {
            "content_hash": content_hash,
            "vault_path": str(vault_path.resolve()),
            "byte_size": len(encoded),
            "source_uri": source_uri,
            "metadata": metadata,
        }
