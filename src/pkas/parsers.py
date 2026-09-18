import csv
import io
import json
import mimetypes
import posixpath
import re
import shutil
import struct
import subprocess
import xml.etree.ElementTree as ElementTree
import zipfile
import zlib
from collections import Counter
from dataclasses import dataclass, field
from email import policy
from email.parser import BytesParser
from itertools import zip_longest
from pathlib import Path
from typing import Any, cast

TEXT_EXTENSIONS = {
    ".bat",
    ".c",
    ".cfg",
    ".cpp",
    ".cs",
    ".css",
    ".go",
    ".h",
    ".hpp",
    ".ini",
    ".java",
    ".js",
    ".jsx",
    ".kt",
    ".log",
    ".md",
    ".ps1",
    ".py",
    ".rs",
    ".scss",
    ".sh",
    ".sql",
    ".svelte",
    ".swift",
    ".toml",
    ".ts",
    ".tsx",
    ".txt",
    ".vue",
    ".xml",
    ".yaml",
    ".yml",
}
OOXML_WORD_EXTENSIONS = {".docx", ".docm", ".dotx", ".dotm"}
OOXML_SHEET_EXTENSIONS = {".xlsx", ".xlsm", ".xltx", ".xltm"}
OOXML_PRESENTATION_EXTENSIONS = {
    ".pptx",
    ".pptm",
    ".potx",
    ".potm",
    ".ppsx",
    ".ppsm",
}

NATIVE_DOCUMENT_EXTENSIONS = {
    ".csv",
    ".tsv",
    ".eml",
    ".htm",
    ".html",
    ".json",
    ".jsonl",
    ".odp",
    ".ods",
    ".odt",
    ".pdf",
    ".rtf",
} | OOXML_WORD_EXTENSIONS | OOXML_SHEET_EXTENSIONS | OOXML_PRESENTATION_EXTENSIONS
VISUAL_EXTENSIONS = {
    ".bmp",
    ".jpeg",
    ".jpg",
    ".png",
    ".tif",
    ".tiff",
    ".webp",
}
DOCLING_EXTENSIONS = {
    ".doc",
    ".epub",
    ".latex",
    ".msg",
    ".ppt",
    ".tex",
    ".xls",
}
SUPPORTED_EXTENSIONS = (
    TEXT_EXTENSIONS | NATIVE_DOCUMENT_EXTENSIONS | VISUAL_EXTENSIONS | DOCLING_EXTENSIONS
)


class UnsupportedFormatError(ValueError):
    pass


class ParseError(ValueError):
    pass


@dataclass(slots=True)
class ParsedMessage:
    sequence: int
    text: str
    speaker: str | None = None
    sent_at: str | None = None
    conversation_id: str | None = None
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass(slots=True)
class ParsedBlock:
    kind: str
    text: str
    locator: str
    page: int | None = None
    bbox: tuple[float, float, float, float] | None = None
    metadata: dict[str, Any] = field(default_factory=dict)

    def as_metadata(self) -> dict[str, Any]:
        item: dict[str, Any] = {
            "kind": self.kind,
            "locator": self.locator,
            "char_count": len(self.text),
        }
        if self.page is not None:
            item["page"] = self.page
        if self.bbox is not None:
            item["bbox"] = [round(value, 2) for value in self.bbox]
        if self.metadata:
            item["metadata"] = self.metadata
        return item


@dataclass(slots=True)
class ParsedDocument:
    title: str
    text: str
    parser_name: str
    parser_version: str = "2"
    mime_type: str | None = None
    language: str | None = None
    event_time: str | None = None
    metadata: dict[str, Any] = field(default_factory=dict)
    messages: list[ParsedMessage] = field(default_factory=list)
    blocks: list[ParsedBlock] = field(default_factory=list)


def _decode_text(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "utf-16", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _clean_text(value: str) -> str:
    return re.sub(r"[ \t]+", " ", value.replace("\x00", " ")).strip()


def _line_blocks(text: str, *, prefix: str = "paragraph") -> list[ParsedBlock]:
    blocks: list[ParsedBlock] = []
    for index, raw in enumerate(text.splitlines(), start=1):
        clean = _clean_text(raw)
        if not clean:
            continue
        kind = "heading" if clean.startswith("#") else "paragraph"
        blocks.append(ParsedBlock(kind=kind, text=clean, locator=f"{prefix}:{index}"))
    return blocks


def _document_text(blocks: list[ParsedBlock]) -> str:
    return "\n\n".join(block.text for block in blocks if block.text.strip()).strip()


def _local_image_dimensions(raw: bytes) -> tuple[str, int | None, int | None]:
    """Read common image container geometry without decoding pixels or metadata."""
    if raw.startswith(b"\x89PNG\r\n\x1a\n") and len(raw) >= 24:
        width, height = struct.unpack(">II", raw[16:24])
        return "PNG", width, height
    if raw.startswith((b"GIF87a", b"GIF89a")) and len(raw) >= 10:
        width, height = struct.unpack("<HH", raw[6:10])
        return "GIF", width, height
    if raw.startswith(b"BM") and len(raw) >= 26:
        width, height = struct.unpack("<ii", raw[18:26])
        return "BMP", abs(width), abs(height)
    if raw.startswith(b"RIFF") and raw[8:12] == b"WEBP" and len(raw) >= 30:
        if raw[12:16] == b"VP8X":
            width = int.from_bytes(raw[24:27], "little") + 1
            height = int.from_bytes(raw[27:30], "little") + 1
            return "WebP", width, height
        return "WebP", None, None
    if raw.startswith(b"\xff\xd8"):
        offset = 2
        sof_markers = {
            0xC0,
            0xC1,
            0xC2,
            0xC3,
            0xC5,
            0xC6,
            0xC7,
            0xC9,
            0xCA,
            0xCB,
            0xCD,
            0xCE,
            0xCF,
        }
        while offset + 9 <= len(raw):
            if raw[offset] != 0xFF:
                offset += 1
                continue
            marker_offset = offset + 1
            while marker_offset < len(raw) and raw[marker_offset] == 0xFF:
                marker_offset += 1
            if marker_offset >= len(raw):
                break
            marker = raw[marker_offset]
            offset = marker_offset + 1
            if marker in {0xD8, 0xD9, 0x01} or 0xD0 <= marker <= 0xD7:
                continue
            if offset + 2 > len(raw):
                break
            segment_length = int.from_bytes(raw[offset : offset + 2], "big")
            if segment_length < 2 or offset + segment_length > len(raw):
                break
            if marker in sof_markers and segment_length >= 7:
                height = int.from_bytes(raw[offset + 3 : offset + 5], "big")
                width = int.from_bytes(raw[offset + 5 : offset + 7], "big")
                return "JPEG", width, height
            offset += segment_length
        return "JPEG", None, None
    if raw.startswith((b"II*\x00", b"MM\x00*")):
        return "TIFF", None, None
    return "image", None, None


_IMAGE_DESCRIPTION_KEYS = {
    "title",
    "description",
    "comment",
    "subject",
    "keywords",
    "author",
    "artist",
    "creator",
    "copyright",
    "rights",
    "label",
}
_IMAGE_DESCRIPTION_LIMIT = 2_000


def _clean_image_description(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("\x00", " ")).strip()[:_IMAGE_DESCRIPTION_LIMIT]


def _bounded_zlib_text(raw: bytes) -> str:
    """Decode a PNG compressed text entry without accepting unbounded output."""
    try:
        decoder = zlib.decompressobj()
        decoded = decoder.decompress(raw, _IMAGE_DESCRIPTION_LIMIT + 1)
        if decoder.unconsumed_tail or len(decoded) > _IMAGE_DESCRIPTION_LIMIT:
            return ""
        return _clean_image_description(_decode_text(decoded))
    except zlib.error:
        return ""


def _xmp_descriptions(raw: bytes) -> list[str]:
    """Keep only authored descriptive XMP fields, never location/device metadata."""
    try:
        root = ElementTree.fromstring(raw)
    except (ElementTree.ParseError, UnicodeDecodeError):
        return []
    values: list[str] = []
    for element in root.iter():
        name = element.tag.rsplit("}", 1)[-1].lower()
        if name not in _IMAGE_DESCRIPTION_KEYS:
            continue
        value = _clean_image_description(" ".join(part for part in element.itertext() if part))
        if value and value not in values:
            values.append(value)
    return values[:8]


def _png_descriptions(raw: bytes) -> list[str]:
    if not raw.startswith(b"\x89PNG\r\n\x1a\n"):
        return []
    values: list[str] = []
    offset = 8
    while offset + 12 <= len(raw):
        size = int.from_bytes(raw[offset : offset + 4], "big")
        end = offset + 12 + size
        if size > 65_536 or end > len(raw):
            break
        chunk_type = raw[offset + 4 : offset + 8]
        payload = raw[offset + 8 : offset + 8 + size]
        if chunk_type == b"tEXt" and b"\x00" in payload:
            key, value = payload.split(b"\x00", 1)
            if _decode_text(key).strip().lower() in _IMAGE_DESCRIPTION_KEYS:
                clean = _clean_image_description(_decode_text(value))
                if clean:
                    values.append(clean)
        elif chunk_type == b"zTXt" and b"\x00" in payload:
            key, rest = payload.split(b"\x00", 1)
            if _decode_text(key).strip().lower() in _IMAGE_DESCRIPTION_KEYS and rest[:1] == b"\x00":
                clean = _bounded_zlib_text(rest[1:])
                if clean:
                    values.append(clean)
        elif chunk_type == b"iTXt" and b"\x00" in payload:
            fields = payload.split(b"\x00", 4)
            if len(fields) == 5:
                key, compression_flag, _language, _translated, value = fields
                key_name = _decode_text(key).strip().lower()
                text = (
                    _bounded_zlib_text(value)
                    if compression_flag == b"\x01"
                    else _clean_image_description(_decode_text(value))
                )
                if key_name == "xml:com.adobe.xmp":
                    values.extend(_xmp_descriptions(value))
                elif key_name in _IMAGE_DESCRIPTION_KEYS and text:
                    values.append(text)
        if chunk_type == b"IEND":
            break
        offset = end
    return list(dict.fromkeys(values))[:8]


def _jpeg_xmp_descriptions(raw: bytes) -> list[str]:
    if not raw.startswith(b"\xff\xd8"):
        return []
    offset = 2
    prefix = b"http://ns.adobe.com/xap/1.0/\x00"
    while offset + 4 <= len(raw):
        if raw[offset] != 0xFF:
            offset += 1
            continue
        marker = raw[offset + 1]
        offset += 2
        if marker in {0xD8, 0xD9} or 0xD0 <= marker <= 0xD7:
            continue
        if offset + 2 > len(raw):
            break
        size = int.from_bytes(raw[offset : offset + 2], "big")
        if size < 2 or offset + size > len(raw):
            break
        payload = raw[offset + 2 : offset + size]
        if marker == 0xE1 and payload.startswith(prefix):
            return _xmp_descriptions(payload[len(prefix) :])
        offset += size
    return []


def _local_image_descriptions(raw: bytes) -> list[str]:
    """Read authored local descriptions only; pixels, EXIF GPS and device data stay unread."""
    return _png_descriptions(raw) or _jpeg_xmp_descriptions(raw)


def _parse_visual_file(path: Path) -> ParsedDocument:
    """Create an honest local catalog record for a standalone image.

    It never reads pixels, EXIF GPS, device details, faces, or OCR text. Authored
    PNG/XMP title, description and keyword fields are local document metadata,
    so they are included when present while preserving the visual-quality signal.
    """
    try:
        raw = path.read_bytes()[:262_144]
        image_format, width, height = _local_image_dimensions(raw)
        details = [f"格式 {image_format}"]
        if width and height:
            details.append(f"尺寸 {width} × {height}")
        descriptions = _local_image_descriptions(raw)
        if descriptions:
            details.append("本地嵌入说明 " + "；".join(descriptions))
        details.append("未读取像素文字或图像语义")
        text = "图片资料元信息：" + "；".join(details) + "。"
        blocks = [
            ParsedBlock(
                kind="image-metadata",
                text=text,
                locator="image:metadata:1",
                metadata={"local_only": True, "pixel_content_not_read": True},
            )
        ]
        blocks.extend(
            ParsedBlock(
                kind="image-description",
                text=f"图片本地嵌入说明：{description}",
                locator=f"image:description:{index}",
                metadata={"local_only": True, "pixel_content_not_read": True},
            )
            for index, description in enumerate(descriptions, start=1)
        )
        return ParsedDocument(
            title=path.stem,
            text=text,
            parser_name="local-image-metadata",
            mime_type=mimetypes.guess_type(path.name)[0] or "application/octet-stream",
            blocks=blocks,
            metadata={
                "image_format": image_format,
                "image_width": width,
                "image_height": height,
                "visual_content_status": "metadata_only",
                "pixel_content_not_read": True,
                "local_description_count": len(descriptions),
            },
        )
    except OSError as exc:
        raise ParseError(f"图片元信息读取失败：{exc}") from exc


def _suppress_repeated_pdf_margins(
    blocks: list[ParsedBlock], page_count: int
) -> tuple[list[ParsedBlock], int]:
    if page_count < 3:
        return blocks, 0
    candidates = Counter(
        re.sub(r"\s+", "", block.text).lower()
        for block in blocks
        if block.metadata.get("page_zone") in {"header", "footer"}
        and 1 <= len(re.sub(r"\s+", "", block.text)) <= 120
    )
    threshold = max(3, (page_count * 3 + 4) // 5)
    repeated = {text for text, count in candidates.items() if count >= threshold}
    kept = [
        block
        for block in blocks
        if re.sub(r"\s+", "", block.text).lower() not in repeated
        or block.metadata.get("page_zone") not in {"header", "footer"}
    ]
    return kept, len(blocks) - len(kept)


def _poppler_pdf_text_fallback(
    path: Path,
    *,
    page_numbers: tuple[int, ...] | None = None,
) -> tuple[list[ParsedBlock], str]:
    """Use local Poppler when PyMuPDF could not read selected page text.

    This is text extraction rather than OCR. A scanned PDF remains unreadable
    when neither parser finds embedded text, and no file is uploaded or written.
    Poppler reads the source once; page filtering happens only after extraction,
    so a mixed PDF never creates a second copy of already-native page text.
    """
    executable = shutil.which("pdftotext")
    if not executable:
        return [], "unavailable"
    try:
        completed = subprocess.run(
            [executable, "-layout", str(path), "-"],
            stdin=subprocess.DEVNULL,
            stdout=subprocess.PIPE,
            stderr=subprocess.DEVNULL,
            check=False,
            timeout=60,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
    except (OSError, subprocess.TimeoutExpired):
        return [], "failed"
    if completed.returncode != 0 or not completed.stdout:
        return [], "empty"
    text = _decode_text(completed.stdout)
    selected_pages = set(page_numbers or ())
    blocks: list[ParsedBlock] = []
    for page_number, page_text in enumerate(text.split("\f"), start=1):
        if selected_pages and page_number not in selected_pages:
            continue
        for line_number, raw_line in enumerate(page_text.splitlines(), start=1):
            clean = _clean_text(raw_line)
            if clean:
                blocks.append(
                    ParsedBlock(
                        kind="paragraph",
                        text=clean,
                        locator=f"page:{page_number}/poppler-line:{line_number}",
                        page=page_number,
                        metadata={"local_fallback": "poppler-pdftotext"},
                    )
                )
    return blocks, "used" if blocks else "empty"


def _message_from_mapping(item: dict[str, Any], sequence: int) -> ParsedMessage | None:
    text_value = item.get("text") or item.get("content") or item.get("message") or item.get("body")
    if text_value is None:
        return None
    if not isinstance(text_value, str):
        text_value = json.dumps(text_value, ensure_ascii=False)
    speaker = item.get("speaker") or item.get("sender") or item.get("from") or item.get("role")
    sent_at = item.get("timestamp") or item.get("sent_at") or item.get("time") or item.get("date")
    conversation_id = (
        item.get("conversation_id")
        or item.get("chat_id")
        or item.get("thread_id")
        or item.get("room")
    )
    excluded = {
        "text",
        "content",
        "message",
        "body",
        "speaker",
        "sender",
        "from",
        "role",
        "timestamp",
        "sent_at",
        "time",
        "date",
        "conversation_id",
        "chat_id",
        "thread_id",
        "room",
    }
    metadata = {key: value for key, value in item.items() if key not in excluded}
    return ParsedMessage(
        sequence=sequence,
        text=text_value.strip(),
        speaker=str(speaker) if speaker is not None else None,
        sent_at=str(sent_at) if sent_at is not None else None,
        conversation_id=str(conversation_id) if conversation_id is not None else None,
        metadata=metadata,
    )


def _message_to_text(message: ParsedMessage) -> str:
    prefix_parts = [value for value in (message.sent_at, message.speaker) if value]
    prefix = " · ".join(prefix_parts)
    return f"[{prefix}] {message.text}" if prefix else message.text


def _messages_to_text(messages: list[ParsedMessage]) -> str:
    return "\n".join(_message_to_text(message) for message in messages)


def _parse_json(path: Path, json_lines: bool = False) -> ParsedDocument:
    raw_text = _decode_text(path.read_bytes())
    try:
        if json_lines:
            data: Any = [json.loads(line) for line in raw_text.splitlines() if line.strip()]
        else:
            data = json.loads(raw_text)
    except json.JSONDecodeError:
        # Some configuration exports and tool dumps are labelled JSON despite
        # trailing commas or truncated records. Preserve readable text, but
        # make the missing structural interpretation explicit.
        blocks = _line_blocks(raw_text, prefix="invalid-json-line")
        return ParsedDocument(
            title=path.stem,
            text=raw_text,
            parser_name=(
                "invalid-jsonl-text-fallback"
                if json_lines
                else "invalid-json-text-fallback"
            ),
            mime_type="text/plain",
            blocks=blocks,
            metadata={
                "json_structure_valid": False,
                "json_structure_not_read": True,
                "fallback_reason": "invalid_json",
            },
        )

    candidates: list[Any]
    if isinstance(data, list):
        candidates = data
    elif isinstance(data, dict):
        nested = data.get("messages") or data.get("items") or data.get("records")
        candidates = nested if isinstance(nested, list) else [data]
    else:
        candidates = []

    messages = [
        message
        for index, item in enumerate(candidates)
        if isinstance(item, dict)
        if (message := _message_from_mapping(item, index)) is not None
    ]
    text = (
        _messages_to_text(messages) if messages else json.dumps(data, ensure_ascii=False, indent=2)
    )
    blocks = [
        ParsedBlock(
            kind="message",
            text=_message_to_text(message),
            locator=f"record:{message.sequence}",
        )
        for message in messages
    ] or _line_blocks(text, prefix="json-line")
    return ParsedDocument(
        title=path.stem,
        text=text,
        parser_name="jsonl" if json_lines else "json",
        mime_type="application/x-ndjson" if json_lines else "application/json",
        messages=messages,
        blocks=blocks,
        metadata={"record_count": len(candidates)},
    )


def _parse_delimited_table(
    path: Path,
    *,
    delimiter: str,
    parser_name: str,
    mime_type: str,
) -> ParsedDocument:
    """Parse CSV/TSV without flattening quoted multiline cells.

    Spreadsheet-like exports are common project evidence.  Keep each row as a
    locatable block and mark the first non-empty row as the header so later
    chunking/ranking can retain field context instead of treating the export as
    unrelated plain text lines.
    """
    text = _decode_text(path.read_bytes())
    rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter))
    blocks: list[ParsedBlock] = []
    header_seen = False
    column_count = 0
    for row_index, row in enumerate(rows, start=1):
        cells = [re.sub(r"\s+", " ", cell).strip() for cell in row]
        if not any(cells):
            continue
        column_count = max(column_count, len(cells))
        blocks.append(
            ParsedBlock(
                kind="table-row",
                text=" | ".join(cells),
                locator=f"row:{row_index}",
                metadata={
                    "table_id": "delimited:1",
                    "row_index": row_index,
                    "is_header": not header_seen,
                    "column_count": len(cells),
                },
            )
        )
        header_seen = True
    return ParsedDocument(
        title=path.stem,
        text=_document_text(blocks),
        parser_name=parser_name,
        mime_type=mime_type,
        blocks=blocks,
        metadata={
            "row_count": len(blocks),
            "column_count": column_count,
            "delimiter": "tab" if delimiter == "\t" else "comma",
        },
    )


def _parse_csv(path: Path) -> ParsedDocument:
    return _parse_delimited_table(
        path,
        delimiter=",",
        parser_name="csv-structured",
        mime_type="text/csv",
    )


def _parse_tsv(path: Path) -> ParsedDocument:
    return _parse_delimited_table(
        path,
        delimiter="\t",
        parser_name="tsv-structured",
        mime_type="text/tab-separated-values",
    )


def _parse_html(path: Path) -> ParsedDocument:
    from bs4 import BeautifulSoup

    html = _decode_text(path.read_bytes())
    soup = BeautifulSoup(html, "lxml")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    title = soup.title.get_text(" ", strip=True) if soup.title else path.stem
    blocks: list[ParsedBlock] = []
    index = 0
    for tag in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "tr"]):
        if tag.name in {"p", "li"} and tag.find_parent("tr"):
            continue
        clean = _clean_text(tag.get_text(" ", strip=True))
        if not clean:
            continue
        index += 1
        if tag.name and tag.name.startswith("h"):
            kind = "heading"
        elif tag.name == "tr":
            kind = "table-row"
        else:
            kind = "paragraph"
        blocks.append(ParsedBlock(kind=kind, text=clean, locator=f"html:{index}"))
    if not blocks:
        text = "\n".join(line.strip() for line in soup.get_text("\n").splitlines() if line.strip())
        blocks = _line_blocks(text, prefix="html-line")
    return ParsedDocument(
        title=title,
        text=_document_text(blocks),
        parser_name="beautifulsoup",
        mime_type="text/html",
        blocks=blocks,
    )


def _parse_pdf(path: Path) -> ParsedDocument:
    import fitz

    try:
        document = fitz.open(path)
        metadata = dict(document.metadata or {})
        blocks: list[ParsedBlock] = []
        page_stats: list[dict[str, Any]] = []
        visual_pages: list[int] = []
        annotation_count = 0
        hyperlink_count = 0
        outline_count = 0
        form_field_count = 0
        redacted_form_field_count = 0
        for outline_index, entry in enumerate(document.get_toc(simple=True), start=1):
            if len(entry) < 3:
                continue
            level, title, page_number = entry[:3]
            clean = _clean_text(str(title))
            if not clean:
                continue
            page = int(page_number) if int(page_number) > 0 else None
            blocks.append(
                ParsedBlock(
                    kind="heading" if int(level) <= 2 else "outline",
                    text=clean,
                    locator=f"outline:{outline_index}",
                    page=page,
                    metadata={"outline_level": int(level)},
                )
            )
            outline_count += 1
        for zero_based_page in range(len(document)):
            page_index = zero_based_page + 1
            page = document[zero_based_page]
            page_height = float(page.rect.height)
            page_blocks = cast(list[tuple[Any, ...]], page.get_text("blocks", sort=True))
            page_text_chars = 0
            text_block_count = 0
            for block_index, raw in enumerate(page_blocks, start=1):
                if len(raw) < 7 or int(raw[6]) != 0:
                    continue
                clean = _clean_text(str(raw[4]))
                if not clean:
                    continue
                bbox = tuple(float(value) for value in raw[:4])
                page_zone = (
                    "header"
                    if bbox[1] <= page_height * 0.12
                    else "footer"
                    if bbox[3] >= page_height * 0.88
                    else "body"
                )
                blocks.append(
                    ParsedBlock(
                        kind="paragraph",
                        text=clean,
                        locator=f"page:{page_index}/block:{block_index}",
                        page=page_index,
                        bbox=cast(tuple[float, float, float, float], bbox),
                        metadata={"page_zone": page_zone},
                    )
                )
                page_text_chars += len(clean)
                text_block_count += 1
            try:
                annotations = page.annots()
                for annotation_index, annotation in enumerate(annotations or (), start=1):
                    info = dict(annotation.info or {})
                    fields = [
                        value
                        for value in (
                            _clean_text(str(info.get("title") or "")),
                            _clean_text(str(info.get("subject") or "")),
                            _clean_text(str(info.get("content") or "")),
                        )
                        if value
                    ]
                    if not fields:
                        continue
                    blocks.append(
                        ParsedBlock(
                            kind="pdf-annotation",
                            text="PDF批注：" + " | ".join(dict.fromkeys(fields)),
                            locator=f"page:{page_index}/annotation:{annotation_index}",
                            page=page_index,
                            bbox=cast(
                                tuple[float, float, float, float],
                                tuple(float(value) for value in annotation.rect),
                            ),
                            metadata={"annotation_type": str(annotation.type[1])},
                        )
                    )
                    annotation_count += 1
            except (RuntimeError, ValueError):
                pass
            try:
                for field_index, widget in enumerate(page.widgets() or (), start=1):
                    field_name = _clean_text(str(widget.field_name or ""))
                    field_value = _clean_text(str(widget.field_value or ""))
                    field_type = _clean_text(str(widget.field_type_string or "未声明类型"))
                    field_flags = int(widget.field_flags or 0)
                    is_sensitive = bool(
                        re.search(
                            r"password|passcode|token|secret|key|密码|口令|密钥",
                            field_name,
                            flags=re.IGNORECASE,
                        )
                    ) or bool(field_flags & 8192)
                    if is_sensitive:
                        redacted_form_field_count += 1
                        if field_value:
                            # Filled form widgets can also be returned by the
                            # page text extractor.  Redact the exact value from
                            # every block already collected for this page, not
                            # only from the dedicated form-field block.
                            for block in blocks:
                                if block.page == page_index and field_value in block.text:
                                    block.text = block.text.replace(
                                        field_value, "[已隐藏表单敏感值]"
                                    )
                        continue
                    if not field_name and not field_value:
                        continue
                    details = [f"类型 {field_type}"]
                    if field_name:
                        details.insert(0, f"字段 {field_name}")
                    if field_value:
                        details.append(f"值 {field_value}")
                    blocks.append(
                        ParsedBlock(
                            kind="pdf-form-field",
                            text="PDF 表单：" + "；".join(details),
                            locator=f"page:{page_index}/form-field:{field_index}",
                            page=page_index,
                            bbox=cast(
                                tuple[float, float, float, float],
                                tuple(float(value) for value in widget.rect),
                            ),
                            metadata={
                                "field_name": field_name or None,
                                "field_type": field_type,
                                "field_has_value": bool(field_value),
                                "local_only": True,
                            },
                        )
                    )
                    form_field_count += 1
            except (RuntimeError, ValueError, AttributeError):
                pass
            for link_index, link in enumerate(page.get_links(), start=1):
                uri = _clean_text(str(link.get("uri") or ""))
                if not uri:
                    continue
                blocks.append(
                    ParsedBlock(
                        kind="hyperlink",
                        text=f"链接：{uri}",
                        locator=f"page:{page_index}/link:{link_index}",
                        page=page_index,
                        metadata={"link_kind": int(link.get("kind") or 0)},
                    )
                )
                hyperlink_count += 1
            image_count = len(page.get_images(full=True))
            if page_text_chars < 40 and (image_count > 0 or page_text_chars == 0):
                visual_pages.append(page_index)
            page_stats.append(
                {
                    "page": page_index,
                    "text_chars": page_text_chars,
                    "text_blocks": text_block_count,
                    "images": image_count,
                }
            )
        blocks, suppressed_margin_blocks = _suppress_repeated_pdf_margins(
            blocks, len(document)
        )
        poppler_blocks: list[ParsedBlock] = []
        poppler_fallback = "not_needed"
        poppler_page_fallback = "not_needed"
        missing_native_pages = tuple(
            int(item["page"])
            for item in page_stats
            if int(item["text_chars"]) == 0
        )
        if not any(int(item["text_chars"]) for item in page_stats):
            poppler_blocks, poppler_fallback = _poppler_pdf_text_fallback(path)
            poppler_page_fallback = poppler_fallback
            blocks.extend(poppler_blocks)
        elif missing_native_pages:
            # Do not spend unlimited time on image-heavy PDFs. This is a local
            # text fallback, not OCR; visual parsing remains the explicit path
            # for the remaining pages.
            selected_pages = missing_native_pages[:8]
            poppler_blocks, poppler_page_fallback = _poppler_pdf_text_fallback(
                path,
                page_numbers=selected_pages,
            )
            poppler_fallback = (
                "used_on_missing_native_pages"
                if poppler_blocks
                else poppler_page_fallback
            )
            blocks.extend(poppler_blocks)
            if len(missing_native_pages) > len(selected_pages):
                poppler_page_fallback += "+page_cap_reached"
        title = metadata.get("title") or path.stem
        result = ParsedDocument(
            title=title,
            text=_document_text(blocks),
            parser_name=(
                "pymupdf-blocks+poppler-pdftotext"
                if poppler_blocks
                else "pymupdf-blocks"
            ),
            mime_type="application/pdf",
            blocks=blocks,
            metadata={
                "page_count": len(document),
                "page_stats": page_stats,
                "visual_page_indexes": visual_pages,
                "suppressed_repeated_margin_blocks": suppressed_margin_blocks,
                "outline_count": outline_count,
                "annotation_count": annotation_count,
                "form_field_count": form_field_count,
                "redacted_form_field_count": redacted_form_field_count,
                "hyperlink_count": hyperlink_count,
                "poppler_fallback": poppler_fallback,
                "poppler_page_fallback": poppler_page_fallback,
                "missing_native_text_pages": list(missing_native_pages),
                "pdf_metadata": metadata,
            },
        )
        document.close()
        return result
    except Exception as exc:
        raise ParseError(f"PDF 解析失败：{exc}") from exc


def _supplemental_docx_blocks(path: Path, existing_text: str) -> list[ParsedBlock]:
    from lxml import etree  # pyright: ignore[reportAttributeAccessIssue]

    existing = re.sub(r"\s+", "", existing_text)
    supplemental: list[ParsedBlock] = []
    part_index = 0
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        parts = sorted(
            name
            for name in names
            if re.fullmatch(
                r"word/(?:header\d+|footer\d+|footnotes|endnotes|comments)\.xml",
                name,
            )
        )
        if "word/document.xml" in names:
            root = etree.fromstring(archive.read("word/document.xml"))
            for text_box in root.xpath(".//*[local-name()='txbxContent']"):
                part_index += 1
                text = " ".join(
                    value.strip()
                    for value in text_box.xpath(".//*[local-name()='t']/text()")
                    if value.strip()
                )
                normalized = re.sub(r"\s+", "", text)
                if normalized and normalized not in existing:
                    supplemental.append(
                        ParsedBlock(
                            kind="text-box",
                            text=text,
                            locator=f"docx:text-box:{part_index}",
                        )
                    )
                    existing += normalized
        for part in parts:
            root = etree.fromstring(archive.read(part))
            part_kind = Path(part).stem.rstrip("0123456789")
            for paragraph_index, paragraph in enumerate(
                root.xpath(".//*[local-name()='p']"),
                start=1,
            ):
                text = "".join(paragraph.xpath(".//*[local-name()='t']/text()"))
                clean = _clean_text(text)
                normalized = re.sub(r"\s+", "", clean)
                if clean and normalized not in existing:
                    supplemental.append(
                        ParsedBlock(
                            kind=part_kind,
                            text=clean,
                            locator=f"docx:{Path(part).name}/paragraph:{paragraph_index}",
                        )
                    )
                    existing += normalized
    return supplemental


def _parse_docx(path: Path) -> ParsedDocument:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = Document(str(path))
        blocks: list[ParsedBlock] = []
        paragraph_index = 0
        table_index = 0
        for child in document.element.body.iterchildren():
            if child.tag.endswith("}p"):
                paragraph = Paragraph(child, document)
                clean = _clean_text(paragraph.text)
                if not clean:
                    continue
                paragraph_index += 1
                style_name = str(paragraph.style.name or "") if paragraph.style else ""
                kind = "heading" if style_name.lower().startswith("heading") else "paragraph"
                blocks.append(
                    ParsedBlock(
                        kind=kind,
                        text=clean,
                        locator=f"docx:paragraph:{paragraph_index}",
                        metadata={"style": style_name} if style_name else {},
                    )
                )
            elif child.tag.endswith("}tbl"):
                table_index += 1
                table = Table(child, document)
                for row_index, row in enumerate(table.rows, start=1):
                    cells: list[str] = []
                    for cell in row.cells:
                        value = _clean_text(cell.text)
                        if not cells or value != cells[-1]:
                            cells.append(value)
                    text = " | ".join(cells)
                    if text.strip(" |"):
                        blocks.append(
                            ParsedBlock(
                                kind="table-row",
                                text=text,
                                locator=f"docx:table:{table_index}/row:{row_index}",
                                metadata={
                                    "table_id": f"docx:table:{table_index}",
                                    "row_index": row_index,
                                    "is_header": row_index == 1,
                                },
                            )
                        )
        blocks.extend(_supplemental_docx_blocks(path, _document_text(blocks)))
        with zipfile.ZipFile(path) as archive:
            media_count = sum(1 for name in archive.namelist() if name.startswith("word/media/"))
        image_descriptions = _openxml_image_description_blocks(
            path,
            part_pattern=r"word/(?:document|header\d+|footer\d+)\.xml",
            locator_prefix="docx",
            media_pattern=r"word/media/[^/]+",
        )
        blocks.extend(image_descriptions)
        return ParsedDocument(
            title=path.stem,
            text=_document_text(blocks),
            parser_name="python-docx-structured",
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            blocks=blocks,
            metadata={
                "paragraph_count": paragraph_index,
                "table_count": table_index,
                "embedded_image_count": media_count,
                "embedded_image_description_count": len(image_descriptions),
            },
        )
    except Exception as exc:
        raise ParseError(f"DOCX 解析失败：{exc}") from exc


def _xlsx_comments(path: Path) -> list[str]:
    from lxml import etree  # pyright: ignore[reportAttributeAccessIssue]

    comments: list[str] = []
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not re.fullmatch(r"xl/comments\d*\.xml", name):
                continue
            root = etree.fromstring(archive.read(name))
            for comment in root.xpath(".//*[local-name()='comment']"):
                text = "".join(comment.xpath(".//*[local-name()='t']/text()"))
                clean = _clean_text(text)
                if clean:
                    reference = comment.get("ref") or "unknown"
                    comments.append(f"{reference}: {clean}")
    return comments


def _is_meaningful_image_label(value: str) -> bool:
    """Ignore Office's generated image names while preserving authored labels."""
    clean = _clean_text(value)
    if not clean:
        return False
    return not bool(
        re.fullmatch(
            r"(?:image|picture|photo|graphic|图片|图像|照片|图)[ _-]*\d*",
            clean,
            flags=re.IGNORECASE,
        )
    )


def _openxml_image_description_blocks(
    path: Path,
    *,
    part_pattern: str,
    locator_prefix: str,
    media_pattern: str | None = None,
) -> list[ParsedBlock]:
    """Read authored OpenXML image descriptions without reading image pixels.

    Office stores accessible alternative text alongside drawing XML. It is useful
    searchable evidence and local-only; generic generated labels are excluded.
    Some exported diagrams lack Office alt text but retain a title or description
    in the image container itself, so callers can also provide a media pattern.
    """
    from lxml import etree  # pyright: ignore[reportAttributeAccessIssue]

    blocks: list[ParsedBlock] = []
    seen: set[str] = set()
    try:
        with zipfile.ZipFile(path) as archive:
            part_names = sorted(
                name for name in archive.namelist() if re.fullmatch(part_pattern, name)
            )
            for name in part_names:
                root = etree.fromstring(archive.read(name))
                for element in root.iter():
                    local_name = str(element.tag).rsplit("}", 1)[-1]
                    if local_name not in {"docPr", "cNvPr"}:
                        continue
                    attributes = {
                        str(key).rsplit("}", 1)[-1]: str(value)
                        for key, value in element.attrib.items()
                    }
                    candidates = [
                        attributes.get("descr", ""),
                        attributes.get("title", ""),
                    ]
                    labels = [
                        _clean_text(value)
                        for value in candidates
                        if _is_meaningful_image_label(value)
                    ]
                    if not labels:
                        continue
                    text = "图片说明：" + "；".join(dict.fromkeys(labels))
                    key = re.sub(r"\s+", "", text).casefold()
                    if not key or key in seen:
                        continue
                    seen.add(key)
                    blocks.append(
                        ParsedBlock(
                            kind="image-description",
                            text=text,
                            locator=f"{locator_prefix}:image-description:{len(blocks) + 1}",
                            metadata={"source_part": name, "local_only": True},
                        )
                    )
            if media_pattern:
                media_names = sorted(
                    name for name in archive.namelist() if re.fullmatch(media_pattern, name)
                )
                for name in media_names:
                    raw = archive.read(name)[:262_144]
                    for description in _local_image_descriptions(raw):
                        text = "图片说明：" + description
                        key = re.sub(r"\s+", "", text).casefold()
                        if not key or key in seen:
                            continue
                        seen.add(key)
                        blocks.append(
                            ParsedBlock(
                                kind="image-description",
                                text=text,
                                locator=(
                                    f"{locator_prefix}:embedded-image-metadata:"
                                    f"{len(blocks) + 1}"
                                ),
                                metadata={
                                    "source_part": name,
                                    "local_only": True,
                                    "pixel_content_not_read": True,
                                    "description_origin": "embedded-image-metadata",
                                },
                            )
                        )
    except (OSError, zipfile.BadZipFile, etree.XMLSyntaxError):
        return []
    return blocks


def _xlsx_chart_description_blocks(path: Path) -> list[ParsedBlock]:
    """Index authored Excel chart labels and data references without rendering it."""
    blocks: list[ParsedBlock] = []
    try:
        with zipfile.ZipFile(path) as archive:
            chart_names = sorted(
                name
                for name in archive.namelist()
                if re.fullmatch(r"xl/charts/chart\d+\.xml", name)
            )
            for chart_index, name in enumerate(chart_names, start=1):
                root = ElementTree.fromstring(archive.read(name))
                labels = list(
                    dict.fromkeys(
                        clean
                        for element in root.iter()
                        if element.tag.rsplit("}", 1)[-1] == "t"
                        for clean in [_clean_text(str(element.text or ""))]
                        if clean
                    )
                )[:12]
                formulas = list(
                    dict.fromkeys(
                        clean
                        for element in root.iter()
                        if element.tag.rsplit("}", 1)[-1] == "f"
                        for clean in [_clean_text(str(element.text or ""))]
                        if clean
                    )
                )[:12]
                if not labels and not formulas:
                    continue
                details: list[str] = []
                if labels:
                    details.append("标签 " + "；".join(labels))
                if formulas:
                    details.append("数据引用 " + "；".join(formulas))
                blocks.append(
                    ParsedBlock(
                        kind="chart-description",
                        text="Excel 图表：" + "；".join(details),
                        locator=f"xlsx:chart:{chart_index}",
                        metadata={
                            "source_part": name,
                            "local_only": True,
                            "pixel_content_not_read": True,
                            "label_count": len(labels),
                            "formula_count": len(formulas),
                        },
                    )
                )
    except (ElementTree.ParseError, OSError, zipfile.BadZipFile):
        return []
    return blocks


def _xlsx_defined_name_blocks(path: Path) -> list[ParsedBlock]:
    """Extract named ranges as workbook business vocabulary.

    Named ranges often carry more domain meaning than their cell coordinates
    (for example, a customer list or an approval threshold).  Keep the actual
    formula/reference so a search result remains traceable; do not evaluate it.
    """
    blocks: list[ParsedBlock] = []
    try:
        with zipfile.ZipFile(path) as archive:
            root = ElementTree.fromstring(archive.read("xl/workbook.xml"))
        for element in root.iter():
            if element.tag.rsplit("}", 1)[-1] != "definedName":
                continue
            name = _clean_text(element.attrib.get("name", ""))
            reference = _clean_text("".join(element.itertext()))
            if not name or not reference:
                continue
            hidden = str(element.attrib.get("hidden", "")).lower() in {"1", "true"}
            scope = element.attrib.get("localSheetId")
            details = [f"名称 {name}", f"引用 {reference}"]
            if hidden:
                details.append("隐藏名称")
            if scope is not None:
                details.append(f"工作表范围 {scope}")
            blocks.append(
                ParsedBlock(
                    kind="named-range",
                    text="Excel 命名区域：" + "；".join(details),
                    locator=f"xlsx:named-range:{len(blocks) + 1}",
                    metadata={
                        "name": name,
                        "reference": reference,
                        "hidden": hidden,
                        "local_sheet_id": scope,
                        "local_only": True,
                    },
                )
            )
    except (ElementTree.ParseError, KeyError, OSError, zipfile.BadZipFile):
        return []
    return blocks


def _xlsx_data_validation_blocks(
    path: Path,
    *,
    max_validations: int = 500,
) -> tuple[list[ParsedBlock], bool]:
    """Read Excel validation rules without loading or evaluating workbook logic.

    Dropdown lists and numeric/date validation encode business constraints.  The
    XML route works with the existing streaming cell reader and imposes a hard
    cap so a malformed workbook cannot create an unbounded number of blocks.
    """
    blocks: list[ParsedBlock] = []
    truncated = False
    try:
        with zipfile.ZipFile(path) as archive:
            relationships_root = ElementTree.fromstring(
                archive.read("xl/_rels/workbook.xml.rels")
            )
            relationship_targets = {
                relation.attrib.get("Id", ""): relation.attrib.get("Target", "")
                for relation in relationships_root
                if relation.tag.rsplit("}", 1)[-1] == "Relationship"
            }
            workbook_root = ElementTree.fromstring(archive.read("xl/workbook.xml"))
            sheet_parts: list[tuple[str, str]] = []
            for sheet in workbook_root.iter():
                if sheet.tag.rsplit("}", 1)[-1] != "sheet":
                    continue
                sheet_name = _clean_text(sheet.attrib.get("name", ""))
                relationship_id = next(
                    (
                        value
                        for key, value in sheet.attrib.items()
                        if key.rsplit("}", 1)[-1] == "id"
                    ),
                    "",
                )
                target = relationship_targets.get(relationship_id, "")
                member = posixpath.normpath(
                    target.lstrip("/")
                    if target.lstrip("/").startswith("xl/")
                    else posixpath.join("xl", target.lstrip("/"))
                )
                if sheet_name and target and member in archive.namelist():
                    sheet_parts.append((sheet_name, member))

            for sheet_name, member in sheet_parts:
                root = ElementTree.fromstring(archive.read(member))
                for validation in root.iter():
                    if validation.tag.rsplit("}", 1)[-1] != "dataValidation":
                        continue
                    if len(blocks) >= max_validations:
                        truncated = True
                        break
                    cell_range = _clean_text(validation.attrib.get("sqref", ""))
                    validation_type = _clean_text(validation.attrib.get("type", "any"))
                    operator = _clean_text(validation.attrib.get("operator", ""))
                    formula_values: dict[str, str] = {}
                    for child in validation:
                        local_name = child.tag.rsplit("}", 1)[-1]
                        if local_name in {"formula1", "formula2"}:
                            formula = _clean_text("".join(child.itertext()))
                            if formula:
                                formula_values[local_name] = formula
                    details = [
                        f"工作表 {sheet_name}",
                        f"范围 {cell_range or '未声明'}",
                        f"类型 {validation_type}",
                    ]
                    if operator:
                        details.append(f"条件 {operator}")
                    for key in ("formula1", "formula2"):
                        if formula_values.get(key):
                            details.append(f"{key} {formula_values[key]}")
                    blocks.append(
                        ParsedBlock(
                            kind="data-validation",
                            text="Excel 数据校验：" + "；".join(details),
                            locator=f"sheet:{sheet_name}/validation:{len(blocks) + 1}",
                            metadata={
                                "sheet": sheet_name,
                                "range": cell_range,
                                "validation_type": validation_type,
                                "operator": operator or None,
                                "formula1": formula_values.get("formula1"),
                                "formula2": formula_values.get("formula2"),
                                "local_only": True,
                            },
                        )
                    )
                if truncated:
                    break
    except (ElementTree.ParseError, KeyError, OSError, zipfile.BadZipFile):
        return [], False
    return blocks, truncated


def _xlsx_merged_cell_map(
    path: Path,
    *,
    max_cells: int = 200_000,
) -> tuple[dict[str, dict[tuple[int, int], str]], int, int, bool]:
    """Map only real XLSX merged cells to their anchor coordinates.

    The streaming OpenPyXL reader deliberately omits merge geometry.  Reading
    workbook XML preserves that geometry without loading the complete workbook
    into memory.  A hard cap avoids expanding pathological merged ranges.
    """
    from openpyxl.utils.cell import get_column_letter, range_boundaries

    try:
        with zipfile.ZipFile(path) as archive:
            relationships_root = ElementTree.fromstring(
                archive.read("xl/_rels/workbook.xml.rels")
            )
            relationship_targets = {
                relation.attrib.get("Id", ""): relation.attrib.get("Target", "")
                for relation in relationships_root
                if relation.tag.rsplit("}", 1)[-1] == "Relationship"
            }
            workbook_root = ElementTree.fromstring(archive.read("xl/workbook.xml"))
            sheet_parts: dict[str, str] = {}
            for sheet in workbook_root.iter():
                if sheet.tag.rsplit("}", 1)[-1] != "sheet":
                    continue
                sheet_name = sheet.attrib.get("name", "")
                relationship_id = next(
                    (
                        value
                        for key, value in sheet.attrib.items()
                        if key.rsplit("}", 1)[-1] == "id"
                    ),
                    "",
                )
                target = relationship_targets.get(relationship_id, "")
                normalized_target = target.lstrip("/")
                member = posixpath.normpath(
                    normalized_target
                    if normalized_target.startswith("xl/")
                    else posixpath.join("xl", normalized_target)
                )
                if sheet_name and target and member in archive.namelist():
                    sheet_parts[sheet_name] = member

            result: dict[str, dict[tuple[int, int], str]] = {}
            range_count = 0
            cell_count = 0
            truncated = False
            for sheet_name, member in sheet_parts.items():
                root = ElementTree.fromstring(archive.read(member))
                merged: dict[tuple[int, int], str] = {}
                for element in root.iter():
                    if element.tag.rsplit("}", 1)[-1] != "mergeCell":
                        continue
                    reference = element.attrib.get("ref", "")
                    min_col, min_row, max_col, max_row = range_boundaries(reference)
                    if min_col == max_col and min_row == max_row:
                        continue
                    range_count += 1
                    anchor = f"{get_column_letter(min_col)}{min_row}"
                    for row in range(min_row, max_row + 1):
                        for column in range(min_col, max_col + 1):
                            if row == min_row and column == min_col:
                                continue
                            if cell_count >= max_cells:
                                truncated = True
                                break
                            merged[(row, column)] = anchor
                            cell_count += 1
                        if truncated:
                            break
                    if truncated:
                        break
                if merged:
                    result[sheet_name] = merged
                if truncated:
                    break
            return result, range_count, cell_count, truncated
    except (ElementTree.ParseError, KeyError, OSError, ValueError, zipfile.BadZipFile):
        return {}, 0, 0, False


def _parse_xlsx(path: Path) -> ParsedDocument:
    from openpyxl import load_workbook
    from openpyxl.utils.cell import get_column_letter

    try:
        values_book = load_workbook(path, read_only=True, data_only=True)
        formulas_book = load_workbook(path, read_only=True, data_only=False)
        blocks: list[ParsedBlock] = []
        sheet_rows: dict[str, int] = {}
        sheet_states: dict[str, str] = {}
        formula_count = 0
        table_count = 0
        merged_cells, merged_range_count, merged_cell_count, merged_cell_truncated = (
            _xlsx_merged_cell_map(path)
        )
        for value_sheet in values_book.worksheets:
            formula_sheet = formulas_book[value_sheet.title]
            sheet_title = value_sheet.title
            sheet_states[sheet_title] = str(value_sheet.sheet_state)
            sheet_table_count = 0
            blocks.append(
                ParsedBlock(
                    kind="heading",
                    text=f"工作表：{sheet_title}",
                    locator=f"sheet:{sheet_title}",
                    metadata={"sheet_state": str(value_sheet.sheet_state)},
                )
            )
            count = 0
            group: list[dict[str, Any]] = []
            merged_values: dict[str, tuple[Any, Any]] = {}

            def append_group(rows: list[dict[str, Any]], *, sheet_name: str = sheet_title) -> None:
                """Emit one logical table, preserving a cover title when it is clear."""
                nonlocal table_count, sheet_table_count
                if not rows:
                    return
                table_count += 1
                sheet_table_count += 1
                table_id = f"sheet:{sheet_name}"
                if sheet_table_count > 1:
                    table_id = f"{table_id}/table:{sheet_table_count}"

                start = 0
                # A common workbook shape is: one-cell title, multi-column header,
                # then data.  Require three rows before making that inference so a
                # compact one-column table is not silently reclassified as a title.
                if (
                    len(rows) >= 3
                    and rows[0]["nonempty_count"] == 1
                    and rows[1]["nonempty_count"] >= 2
                    and rows[2]["nonempty_count"] >= 2
                ):
                    title = rows[0]
                    blocks.append(
                        ParsedBlock(
                            kind="table-title",
                            text=title["text"],
                            locator=f"sheet:{sheet_name}/row:{title['row_index']}",
                            metadata={
                                "table_id": table_id,
                                "table_index": sheet_table_count,
                                "role": "title",
                            },
                        )
                    )
                    start = 1

                header = rows[start]
                header_labels = list(header["labels"])
                header_locator = f"sheet:{sheet_name}/row:{header['row_index']}"
                for offset, row in enumerate(rows[start:]):
                    is_header = offset == 0
                    row_metadata = {
                        "table_id": table_id,
                        "table_index": sheet_table_count,
                        "row_index": row["row_index"],
                        "is_header": is_header,
                        "column_count": row["nonempty_count"],
                        "header_locator": header_locator,
                        "column_headers": header_labels,
                    }
                    if row["merged_from"]:
                        row_metadata["merged_from"] = row["merged_from"]
                    blocks.append(
                        ParsedBlock(
                            kind="table-row",
                            text=row["text"],
                            locator=f"sheet:{sheet_name}/row:{row['row_index']}",
                            metadata=row_metadata,
                        )
                    )

            value_rows = value_sheet.iter_rows()
            formula_rows = formula_sheet.iter_rows()
            for row_index, pair in enumerate(
                zip_longest(value_rows, formula_rows, fillvalue=()),
                start=1,
            ):
                value_row, formula_row = pair
                rendered: list[str] = []
                labels: list[str] = []
                merged_from_cells: list[str] = []
                for column_index, (value_cell, formula_cell) in enumerate(
                    zip_longest(value_row, formula_row, fillvalue=None),
                    start=1,
                ):
                    cell = formula_cell or value_cell
                    if cell is None:
                        continue
                    coordinate = getattr(
                        cell,
                        "coordinate",
                        f"{get_column_letter(column_index)}{row_index}",
                    )
                    value = value_cell.value if value_cell is not None else None
                    formula = formula_cell.value if formula_cell is not None else None
                    merged_from = merged_cells.get(sheet_title, {}).get(
                        (row_index, column_index)
                    )
                    if value is None and formula is None and merged_from:
                        value, formula = merged_values.get(merged_from, (None, None))
                        if value is not None or formula is not None:
                            merged_from_cells.append(f"{coordinate}←{merged_from}")
                    if value is None and formula is None:
                        continue
                    display = "" if value is None else str(value)
                    if (
                        not merged_from
                        and isinstance(formula, str)
                        and formula.startswith("=")
                    ):
                        formula_count += 1
                        display = f"{display} [公式 {formula}]".strip()
                    rendered.append(f"{coordinate}={display}")
                    if display:
                        labels.append(display)
                    if not merged_from:
                        merged_values[coordinate] = (value, formula)
                if rendered:
                    group.append(
                        {
                            "row_index": row_index,
                            "text": " | ".join(rendered),
                            "labels": labels,
                            "nonempty_count": len(rendered),
                            "merged_from": merged_from_cells,
                        }
                    )
                    count += 1
                elif group:
                    append_group(group)
                    group = []
            append_group(group)
            sheet_rows[sheet_title] = count
        for index, comment in enumerate(_xlsx_comments(path), start=1):
            blocks.append(
                ParsedBlock(kind="comment", text=comment, locator=f"xlsx:comment:{index}")
            )
        image_descriptions = _openxml_image_description_blocks(
            path,
            part_pattern=r"xl/drawings/[^/]+\.xml",
            locator_prefix="xlsx",
            media_pattern=r"xl/media/[^/]+",
        )
        blocks.extend(image_descriptions)
        chart_descriptions = _xlsx_chart_description_blocks(path)
        blocks.extend(chart_descriptions)
        named_ranges = _xlsx_defined_name_blocks(path)
        blocks.extend(named_ranges)
        data_validations, data_validations_truncated = _xlsx_data_validation_blocks(path)
        blocks.extend(data_validations)
        values_book.close()
        formulas_book.close()
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
            image_count = sum(1 for name in names if name.startswith("xl/media/"))
            chart_count = sum(1 for name in names if name.startswith("xl/charts/chart"))
        return ParsedDocument(
            title=path.stem,
            text=_document_text(blocks),
            parser_name="openpyxl-structured",
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            blocks=blocks,
            metadata={
                "sheet_rows": sheet_rows,
                "sheet_states": sheet_states,
                "table_count": table_count,
                "formula_count": formula_count,
                "merged_range_count": merged_range_count,
                "merged_cell_count": merged_cell_count,
                "merged_cell_mapping_truncated": merged_cell_truncated,
                "embedded_image_count": image_count,
                "embedded_image_description_count": len(image_descriptions),
                "chart_count": chart_count,
                "chart_description_count": len(chart_descriptions),
                "named_range_count": len(named_ranges),
                "data_validation_count": len(data_validations),
                "data_validation_truncated": data_validations_truncated,
            },
        )
    except Exception as exc:
        raise ParseError(f"XLSX 解析失败：{exc}") from exc


def _office_xml_paragraphs(raw: bytes) -> list[tuple[str, bool]]:
    from lxml import etree  # pyright: ignore[reportAttributeAccessIssue]

    root = etree.fromstring(raw)
    paragraphs: list[tuple[str, bool]] = []
    for paragraph in root.xpath(".//*[local-name()='p']"):
        values = paragraph.xpath(".//*[local-name()='t']/text()")
        clean = _clean_text("".join(values))
        if not clean:
            continue
        in_table = any(ancestor.tag.endswith("}tc") for ancestor in paragraph.iterancestors())
        paragraphs.append((clean, in_table))
    return paragraphs


def _parse_pptx(path: Path) -> ParsedDocument:
    try:
        blocks: list[ParsedBlock] = []
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()

            def numbered_part_index(name: str) -> int:
                match = re.search(r"(\d+)", Path(name).stem)
                return int(match.group(1)) if match else 0

            slide_names = sorted(
                (name for name in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)),
                key=numbered_part_index,
            )
            for slide_index, name in enumerate(slide_names, start=1):
                paragraphs = _office_xml_paragraphs(archive.read(name))
                for paragraph_index, (text, in_table) in enumerate(paragraphs, start=1):
                    kind = (
                        "table-row"
                        if in_table
                        else ("heading" if paragraph_index == 1 else "paragraph")
                    )
                    blocks.append(
                        ParsedBlock(
                            kind=kind,
                            text=text,
                            locator=f"slide:{slide_index}/block:{paragraph_index}",
                            page=slide_index,
                        )
                    )
            note_names = sorted(
                name for name in names if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", name)
            )
            for note_index, name in enumerate(note_names, start=1):
                for paragraph_index, (text, _in_table) in enumerate(
                    _office_xml_paragraphs(archive.read(name)),
                    start=1,
                ):
                    blocks.append(
                        ParsedBlock(
                            kind="speaker-note",
                            text=text,
                            locator=f"slide-note:{note_index}/paragraph:{paragraph_index}",
                            page=note_index,
                        )
                    )
            image_count = sum(1 for name in names if name.startswith("ppt/media/"))
        image_descriptions = _openxml_image_description_blocks(
            path,
            part_pattern=r"ppt/slides/slide\d+\.xml",
            locator_prefix="pptx",
            media_pattern=r"ppt/media/[^/]+",
        )
        blocks.extend(image_descriptions)
        return ParsedDocument(
            title=path.stem,
            text=_document_text(blocks),
            parser_name="openxml-pptx-structured",
            mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            blocks=blocks,
            metadata={
                "slide_count": len(slide_names),
                "speaker_note_count": len(note_names),
                "embedded_image_count": image_count,
                "embedded_image_description_count": len(image_descriptions),
            },
        )
    except Exception as exc:
        raise ParseError(f"PPTX 解析失败：{exc}") from exc


def _parse_eml(path: Path) -> ParsedDocument:
    try:
        message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        blocks: list[ParsedBlock] = []
        subject = str(message.get("subject") or path.stem)
        for header in ("from", "to", "cc", "date", "subject"):
            value = message.get(header)
            if value:
                blocks.append(
                    ParsedBlock(
                        kind="email-header",
                        text=f"{header.title()}: {value}",
                        locator=f"email:header:{header}",
                    )
                )
        attachments: list[str] = []
        body_index = 0
        for part in message.walk() if message.is_multipart() else [message]:
            disposition = part.get_content_disposition()
            if disposition == "attachment":
                attachments.append(part.get_filename() or "unnamed")
                continue
            if part.get_content_type() not in {"text/plain", "text/html"}:
                continue
            content = part.get_content()
            if not isinstance(content, str):
                continue
            if part.get_content_type() == "text/html":
                from bs4 import BeautifulSoup

                content = BeautifulSoup(content, "lxml").get_text("\n")
            for block in _line_blocks(content, prefix=f"email:body:{body_index}"):
                blocks.append(block)
            body_index += 1
        return ParsedDocument(
            title=subject,
            text=_document_text(blocks),
            parser_name="email-stdlib",
            mime_type="message/rfc822",
            event_time=str(message.get("date")) if message.get("date") else None,
            blocks=blocks,
            metadata={
                "attachment_count": len(attachments),
                "attachment_names": attachments,
            },
        )
    except Exception as exc:
        raise ParseError(f"EML 解析失败：{exc}") from exc


def _rtf_encoding(raw: bytes) -> str:
    """Choose the declared RTF ANSI code page without involving Office/COM."""
    match = re.search(br"\\ansicpg(\d+)", raw)
    codepage = match.group(1).decode("ascii") if match else "1252"
    aliases = {"936": "gbk", "950": "big5", "932": "cp932", "65001": "utf-8"}
    return aliases.get(codepage, f"cp{codepage}")


def _parse_rtf(path: Path) -> ParsedDocument:
    r"""Extract a conservative plain-text view of an RTF file locally.

    RTF is a control language, not an ordinary text file: blindly stripping
    backslashes loses Chinese \u escapes and joins paragraphs.  This compact
    reader keeps text, paragraph/tab separators and declared ANSI bytes while
    deliberately ignoring formatting instructions and embedded objects.
    """
    try:
        raw = path.read_bytes()
        encoding = _rtf_encoding(raw)
        output: list[str] = []
        pending = bytearray()

        def flush() -> None:
            if pending:
                output.append(pending.decode(encoding, errors="replace"))
                pending.clear()

        index = 0
        while index < len(raw):
            current = raw[index]
            if current in {ord("{"), ord("}")}:
                index += 1
                continue
            if current != ord("\\"):
                pending.append(current)
                index += 1
                continue
            index += 1
            if index >= len(raw):
                break
            control = raw[index]
            if control in {ord("\\"), ord("{"), ord("}")}:
                pending.append(control)
                index += 1
                continue
            if control == ord("'") and index + 2 < len(raw):
                try:
                    pending.append(int(raw[index + 1:index + 3].decode("ascii"), 16))
                    index += 3
                    continue
                except ValueError:
                    pass
            if control == ord("u"):
                cursor = index + 1
                sign = 1
                if cursor < len(raw) and raw[cursor] == ord("-"):
                    sign = -1
                    cursor += 1
                start = cursor
                while cursor < len(raw) and chr(raw[cursor]).isdigit():
                    cursor += 1
                if cursor > start:
                    flush()
                    value = sign * int(raw[start:cursor].decode("ascii"))
                    output.append(chr(value if value >= 0 else value + 65536))
                    # RTF commonly includes one ANSI fallback character after \uN.
                    if cursor < len(raw) and raw[cursor] not in {ord("\\"), ord("{"), ord("}")}:
                        cursor += 1
                    index = cursor
                    continue
            if chr(control).isalpha():
                start = index
                while index < len(raw) and chr(raw[index]).isalpha():
                    index += 1
                word = raw[start:index].decode("ascii", errors="ignore").lower()
                if index < len(raw) and raw[index] in {ord("-"), *range(ord("0"), ord("9") + 1)}:
                    index += 1
                    while index < len(raw) and chr(raw[index]).isdigit():
                        index += 1
                if index < len(raw) and raw[index] == ord(" "):
                    index += 1
                if word in {"par", "line"}:
                    flush()
                    output.append("\n")
                elif word == "tab":
                    flush()
                    output.append("\t")
                continue
            if control == ord("~"):
                pending.append(ord(" "))
            elif control == ord("-"):
                pending.append(ord("-"))
            index += 1
        flush()
        text = _clean_text("".join(output).replace("\r", "\n"))
        blocks = _line_blocks(text, prefix="rtf")
        return ParsedDocument(
            title=path.stem,
            text=text,
            parser_name="rtf-local-text",
            mime_type="application/rtf",
            blocks=blocks,
            metadata={"encoding": encoding, "embedded_objects_read": False},
        )
    except Exception as exc:
        raise ParseError(f"RTF 解析失败：{exc}") from exc


def _odf_local_name(value: str) -> str:
    return value.rsplit("}", 1)[-1]


def _odf_attribute(element: Any, name: str, default: str = "") -> str:
    for key, value in element.attrib.items():
        if _odf_local_name(str(key)) == name:
            return str(value)
    return default


def _odf_text(element: Any) -> str:
    return _clean_text(" ".join(piece.strip() for piece in element.itertext() if piece.strip()))


def _odf_table_blocks(table: Any, table_index: int, *, prefix: str) -> list[ParsedBlock]:
    table_id = f"{prefix}:table:{table_index}"
    table_name = _odf_attribute(table, "name", f"表格 {table_index}")
    blocks = [
        ParsedBlock(
            kind="heading",
            text=f"{table_name}",
            locator=table_id,
            metadata={"table_id": table_id},
        )
    ]
    row_index = 0
    for row in (item for item in table.iter() if _odf_local_name(str(item.tag)) == "table-row"):
        row_index += 1
        values: list[str] = []
        for cell in row:
            if _odf_local_name(str(cell.tag)) not in {"table-cell", "covered-table-cell"}:
                continue
            value = _odf_text(cell)
            formula = _odf_attribute(cell, "formula")
            if formula:
                value = f"{value} [公式 {formula}]".strip()
            repeated_columns = int(_odf_attribute(cell, "number-columns-repeated", "1") or "1")
            if not value:
                continue
            values.append(value if repeated_columns == 1 else f"{value} ×{repeated_columns}")
        if not values:
            continue
        repeated_rows = int(_odf_attribute(row, "number-rows-repeated", "1") or "1")
        in_header = any(
            _odf_local_name(str(ancestor.tag)) == "table-header-rows"
            for ancestor in row.iterancestors()
        )
        text = " | ".join(values)
        if repeated_rows > 1:
            text = f"{text} [重复 {repeated_rows} 行]"
        blocks.append(
            ParsedBlock(
                kind="table-row",
                text=text,
                locator=f"{table_id}/row:{row_index}",
                metadata={
                    "table_id": table_id,
                    "table_name": table_name,
                    "row_index": row_index,
                    "is_header": in_header or row_index == 1,
                    "repeat_count": repeated_rows,
                },
            )
        )
    return blocks


def _odf_content_root(path: Path) -> Any:
    from lxml import etree  # pyright: ignore[reportAttributeAccessIssue]

    with zipfile.ZipFile(path) as archive:
        return etree.fromstring(archive.read("content.xml"))


def _parse_odt(path: Path) -> ParsedDocument:
    try:
        root = _odf_content_root(path)
        body = next(
            (item for item in root.iter() if _odf_local_name(str(item.tag)) == "body"),
            root,
        )
        blocks: list[ParsedBlock] = []
        table_count = 0

        def walk(node: Any) -> None:
            nonlocal table_count
            kind = _odf_local_name(str(node.tag))
            if kind == "table":
                table_count += 1
                blocks.extend(_odf_table_blocks(node, table_count, prefix="odt"))
                return
            if kind in {"h", "p"}:
                text = _odf_text(node)
                if text:
                    blocks.append(
                        ParsedBlock(
                            kind="heading" if kind == "h" else "paragraph",
                            text=text,
                            locator=f"odt:{kind}:{len(blocks) + 1}",
                        )
                    )
                return
            for child in node:
                walk(child)

        walk(body)
        return ParsedDocument(
            title=path.stem,
            text=_document_text(blocks),
            parser_name="opendocument-text-structured",
            mime_type="application/vnd.oasis.opendocument.text",
            blocks=blocks,
            metadata={"table_count": table_count, "embedded_objects_read": False},
        )
    except Exception as exc:
        raise ParseError(f"ODT 解析失败：{exc}") from exc


def _parse_ods(path: Path) -> ParsedDocument:
    try:
        root = _odf_content_root(path)
        sheets = [item for item in root.iter() if _odf_local_name(str(item.tag)) == "table"]
        blocks: list[ParsedBlock] = []
        for sheet_index, sheet in enumerate(sheets, start=1):
            name = _odf_attribute(sheet, "name", f"工作表 {sheet_index}")
            blocks.append(
                ParsedBlock(kind="heading", text=f"工作表：{name}", locator=f"sheet:{sheet_index}")
            )
            blocks.extend(_odf_table_blocks(sheet, sheet_index, prefix="ods"))
        return ParsedDocument(
            title=path.stem,
            text=_document_text(blocks),
            parser_name="opendocument-sheet-structured",
            mime_type="application/vnd.oasis.opendocument.spreadsheet",
            blocks=blocks,
            metadata={"sheet_count": len(sheets), "embedded_objects_read": False},
        )
    except Exception as exc:
        raise ParseError(f"ODS 解析失败：{exc}") from exc


def _parse_odp(path: Path) -> ParsedDocument:
    try:
        root = _odf_content_root(path)
        pages = [item for item in root.iter() if _odf_local_name(str(item.tag)) == "page"]
        blocks: list[ParsedBlock] = []
        for page_index, page in enumerate(pages, start=1):
            paragraphs = [
                _odf_text(item)
                for item in page.iter()
                if _odf_local_name(str(item.tag)) in {"h", "p"}
            ]
            for paragraph_index, text in enumerate((item for item in paragraphs if item), start=1):
                blocks.append(
                    ParsedBlock(
                        kind="heading" if paragraph_index == 1 else "paragraph",
                        text=text,
                        locator=f"slide:{page_index}/block:{paragraph_index}",
                        page=page_index,
                    )
                )
        return ParsedDocument(
            title=path.stem,
            text=_document_text(blocks),
            parser_name="opendocument-presentation-structured",
            mime_type="application/vnd.oasis.opendocument.presentation",
            blocks=blocks,
            metadata={"slide_count": len(pages), "embedded_objects_read": False},
        )
    except Exception as exc:
        raise ParseError(f"ODP 解析失败：{exc}") from exc


def parse_native_file(path: Path) -> ParsedDocument:
    extension = path.suffix.lower()
    if extension in TEXT_EXTENSIONS:
        text = _decode_text(path.read_bytes())
        mime_type = mimetypes.guess_type(path.name)[0] or "text/plain"
        blocks = _line_blocks(text)
        return ParsedDocument(
            title=path.stem,
            text=text,
            parser_name="plain-text",
            mime_type=mime_type,
            blocks=blocks,
        )
    if extension in VISUAL_EXTENSIONS:
        return _parse_visual_file(path)
    if extension == ".csv":
        return _parse_csv(path)
    if extension == ".tsv":
        return _parse_tsv(path)
    if extension == ".json":
        return _parse_json(path)
    if extension == ".jsonl":
        return _parse_json(path, json_lines=True)
    if extension in {".html", ".htm"}:
        return _parse_html(path)
    if extension == ".pdf":
        return _parse_pdf(path)
    if extension == ".rtf":
        return _parse_rtf(path)
    if extension in OOXML_WORD_EXTENSIONS:
        return _parse_docx(path)
    if extension in OOXML_SHEET_EXTENSIONS:
        return _parse_xlsx(path)
    if extension in OOXML_PRESENTATION_EXTENSIONS:
        return _parse_pptx(path)
    if extension == ".odt":
        return _parse_odt(path)
    if extension == ".ods":
        return _parse_ods(path)
    if extension == ".odp":
        return _parse_odp(path)
    if extension == ".eml":
        return _parse_eml(path)
    raise UnsupportedFormatError(f"原生解析器不支持该格式：{extension or '无扩展名'}")


def parse_file(
    path: Path,
    *,
    settings: Any | None = None,
    privacy: str = "private",
    paddle_transport: Any | None = None,
) -> ParsedDocument:
    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFormatError(f"暂不支持该格式：{path.suffix.lower() or '无扩展名'}")
    from pkas.document_extraction import DocumentExtractionPipeline

    return DocumentExtractionPipeline(
        settings=settings,
        paddle_transport=paddle_transport,
    ).parse(path, privacy=privacy)
