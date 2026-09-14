import argparse
import json
import re
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from pkas.config import Settings, get_settings
from pkas.parsers import ParsedDocument, parse_file, parse_native_file


def _normalize(value: str) -> str:
    return re.sub(r"\s+", "", value).casefold()


def text_metrics(expected: str, actual: str) -> dict[str, float]:
    expected_chars = Counter(_normalize(expected))
    actual_chars = Counter(_normalize(actual))
    overlap = sum((expected_chars & actual_chars).values())
    expected_total = sum(expected_chars.values())
    actual_total = sum(actual_chars.values())
    recall = overlap / expected_total if expected_total else 1.0
    precision = overlap / actual_total if actual_total else (1.0 if not expected_total else 0.0)
    f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
    return {
        "text_recall": round(recall, 6),
        "text_precision": round(precision, 6),
        "text_f1": round(f1, 6),
    }


def _legacy_parse(path: Path) -> ParsedDocument | None:
    extension = path.suffix.lower()
    if extension == ".pdf":
        import fitz

        document = fitz.open(path)
        text = "\n\n".join(str(page.get_text("text")) for page in document)
        document.close()
        return ParsedDocument(title=path.stem, text=text, parser_name="legacy-pymupdf")
    if extension == ".docx":
        from docx import Document

        document = Document(str(path))
        lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        return ParsedDocument(title=path.stem, text="\n".join(lines), parser_name="legacy-docx")
    if extension == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        lines: list[str] = []
        for sheet in workbook.worksheets:
            lines.append(f"## {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(value.strip() for value in values):
                    lines.append(" | ".join(values))
        workbook.close()
        return ParsedDocument(title=path.stem, text="\n".join(lines), parser_name="legacy-xlsx")
    try:
        return parse_native_file(path)
    except (ValueError, OSError):
        return None


def _field_recall(expected_fields: list[str], actual: str) -> float:
    if not expected_fields:
        return 1.0
    normalized_actual = _normalize(actual)
    hits = sum(1 for field in expected_fields if _normalize(field) in normalized_actual)
    return round(hits / len(expected_fields), 6)


def _block_type_recall(expected_types: list[str], parsed: ParsedDocument | None) -> float:
    if not expected_types:
        return 1.0
    actual_types = {block.kind.casefold() for block in parsed.blocks} if parsed else set()
    hits = sum(1 for kind in expected_types if kind.casefold() in actual_types)
    return round(hits / len(expected_types), 6)


def evaluate_manifest(
    manifest_path: Path,
    *,
    settings: Settings | None = None,
    privacy: str = "private",
) -> dict[str, Any]:
    active_settings = settings or get_settings()
    records = [
        json.loads(line)
        for line in manifest_path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    profiles: dict[str, list[dict[str, float]]] = {"legacy": [], "hybrid": []}
    errors: list[dict[str, Any]] = []
    for index, record in enumerate(records):
        path = Path(str(record["path"])).expanduser().resolve(strict=True)
        expected_text = str(record.get("expected_text") or "")
        expected_fields = [str(item) for item in record.get("expected_fields", [])]
        expected_types = [str(item) for item in record.get("expected_block_types", [])]
        try:
            legacy = _legacy_parse(path)
            hybrid = parse_file(path, settings=active_settings, privacy=privacy)
        except Exception as exc:
            errors.append({"record": index, "error_type": type(exc).__name__})
            continue
        for profile, parsed in (("legacy", legacy), ("hybrid", hybrid)):
            actual = parsed.text if parsed else ""
            metrics = text_metrics(expected_text, actual)
            metrics["field_recall"] = _field_recall(expected_fields, actual)
            metrics["block_type_recall"] = _block_type_recall(expected_types, parsed)
            profiles[profile].append(metrics)

    summary: dict[str, Any] = {}
    for profile, rows in profiles.items():
        if not rows:
            summary[profile] = {"evaluated_documents": 0}
            continue
        keys = rows[0].keys()
        summary[profile] = {
            "evaluated_documents": len(rows),
            **{key: round(sum(row[key] for row in rows) / len(rows), 6) for key in keys},
        }
    improvements: dict[str, float] = {}
    if profiles["legacy"] and profiles["hybrid"]:
        for key in profiles["legacy"][0]:
            improvements[key] = round(summary["hybrid"][key] - summary["legacy"][key], 6)
    return {
        "status": "completed" if not errors else "warning",
        "manifest_records": len(records),
        "profiles": summary,
        "improvements": improvements,
        "error_count": len(errors),
        "errors": errors,
        "generated_at": datetime.now(UTC).isoformat(),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="评测 PKAS 旧版与混合文档解析能力")
    parser.add_argument("manifest", type=Path, help="UTF-8 JSONL 标注清单")
    parser.add_argument("--output", type=Path)
    parser.add_argument("--privacy", choices=["public", "private", "restricted"], default="private")
    args = parser.parse_args()
    settings = get_settings()
    report = evaluate_manifest(args.manifest, settings=settings, privacy=args.privacy)
    output = args.output or (
        settings.data_root
        / "runs"
        / f"extraction-eval-{datetime.now(UTC).strftime('%Y%m%dT%H%M%SZ')}.json"
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"status": report["status"], "report": str(output)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
